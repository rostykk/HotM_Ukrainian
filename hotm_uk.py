#!/usr/bin/env python3
"""Українська локалізація Heart of the Machine.

Команди:
  sync    витягнути англійський текст з XML гри і оновити робочий файл uk_translation.xlsx
          (нові / змінені / видалені рядки + звіт у changes/)
  build   зібрати GameData/LocalizationImports/Ukrainian.xlsx з робочого файлу
          і повернути мову в LocalizationTypes.xml, якщо оновлення гри її стерло
  status  коротка статистика перекладу
  dump    оновити текстову копію text/*.json (для git) з uk_translation.xlsx і glossary.xlsx
  pseudo  тестовий Ukrainian.xlsx ("ЇҐ·" + англійський текст) для перевірки шрифтів

  glossary   оновити glossary.xlsx кандидатами в терміни
  export     нарізати неперекладені рядки (або терміни глосарію) на JSON-батчі для ШІ
  claim      взяти N батчів з початку (або з кінця, --from-end) і позначити, хто їх перекладає
  batches    стан батчів
  import     імпортувати відповіді ШІ з batches/<раунд>/out/*.json
  terms      редакторська перевірка: де переклад не відповідає глосарію
"""
import argparse
import collections
import datetime
import glob
import json
import os
import re
import shutil
import sys
import zipfile
import xml.etree.ElementTree as ET

try:
    import openpyxl
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    sys.exit("Немає openpyxl. Запускайте через .venv\\Scripts\\python.exe "
             "(створення середовища — див. «Середовище» в README.md)")

HERE = os.path.dirname(os.path.abspath(__file__))
LANG_ID = "Ukrainian"
WORK_FILE = os.path.join(HERE, "uk_translation.xlsx")
LOCTYPE_BLOCK_FILE = os.path.join(HERE, "localization_type_uk.xml")
REFERENCE_LANGS = ["German"]  # довідкові колонки з офіційних перекладів
OFFICIAL_LANGS = ["ChineseSimplified", "ChineseTraditional", "French", "German",
                  "Japanese", "Korean", "PortugueseBrazil", "Spanish"]

# Папки, у яких назва таблиці в ключі відрізняється від назви папки (таблиці ArcenUniversal).
TABLE_NAME_OVERRIDES = {"MusicTracks": "MusicTrack", "9_InputAction": "InputActionTypeData"}

SHEET_MAIN, SHEET_REMOVED, SHEET_INFO = "Переклад", "Видалено", "Інфо"
COL_KEY, COL_EN, COL_UK, COL_STATUS = "Key", "English", "Українська", "Статус"
COL_OLD_EN, COL_NOTE, COL_FILE = "Попередній English", "Примітка для перекладача", "Файл"
STATUSES = ("new", "changed", "auto", "mt", "ok")
COL_BY = "Перекладач"

PLACEHOLDER_RE = re.compile(r"\{[^{}]*\}")
TAG_RE = re.compile(r"</?[a-zA-Z#][^<>]*>")
ILLEGAL_XLSX_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

try:
    # у справжній консолі Windows Python і так пише Unicode; у конвеєрах — примусово UTF-8
    if sys.stdout.isatty():
        sys.stdout.reconfigure(errors="replace")
    else:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass


# ---------------------------------------------------------------- шляхи гри

GAME_ROOT = None  # тека гри з --game/HOTM_DIR; задається в main()


def find_game_dir(arg):
    cands = [arg, os.environ.get("HOTM_DIR"), os.path.dirname(HERE)]
    for c in cands:
        if c and os.path.isdir(os.path.join(c, "GameData", "Configuration")):
            return os.path.abspath(c)
    sys.exit("Не знайдено теку гри. Вкажіть --game \"D:\\...\\HeartOfTheMachine\"")


class Game:
    def __init__(self, root):
        self.root = root
        self.cfg = os.path.join(root, "GameData", "Configuration")
        self.imports = os.path.join(root, "GameData", "LocalizationImports")
        self.expansions = sorted(
            d for d in glob.glob(os.path.join(root, "Expansions", "*")) if os.path.isdir(d))
        self.loctypes = os.path.join(self.cfg, "LocalizationType", "LocalizationTypes.xml")
        self.out_xlsx = os.path.join(self.imports, LANG_ID + ".xlsx")

    def version(self):
        try:
            root = parse_xml_file(os.path.join(self.cfg, "0_GameVersion", "GameVersions.xml"))
            last = [el for el in root if el.tag == "game_version"][-1]
            return last.get("id")
        except Exception:
            return "unknown"

    def build_id(self):
        """Номер збірки Steam з appmanifest_<appid>.acf. Arcen буває змінює тексти, не змінюючи
        id у GameVersions.xml, а buildid змінюється з кожним оновленням. None, якщо не Steam."""
        try:
            appid = open(os.path.join(self.root, "steam_appid.txt"), encoding="utf-8").read().strip()
            acf = os.path.join(os.path.dirname(os.path.dirname(self.root)), f"appmanifest_{appid}.acf")
            m = re.search(r'"buildid"\s+"(\d+)"', open(acf, encoding="utf-8", errors="replace").read())
            return m.group(1) if m else None
        except OSError:
            return None


# ---------------------------------------------------------------- XML

def protect_attr_whitespace(text):
    """Стандартні XML-парсери замінюють переноси рядків і табуляції в атрибутах на пробіли.
    Гра їх зберігає, тому всередині значень атрибутів замінюємо їх на &#10; / &#9;."""
    out = []
    i, n = 0, len(text)
    while i < n:
        for opener, closer in (("<!--", "-->"), ("<![CDATA[", "]]>"), ("<?", "?>")):
            if text.startswith(opener, i):
                j = text.find(closer, i + len(opener))
                j = n if j < 0 else j + len(closer)
                out.append(text[i:j])
                i = j
                break
        else:
            if text[i] != "<":
                out.append(text[i])
                i += 1
                continue
            j, quote, buf = i + 1, None, ["<"]
            while j < n:
                ch = text[j]
                if quote:
                    if ch == quote:
                        quote = None
                        buf.append(ch)
                    elif ch == "\r":
                        if j + 1 < n and text[j + 1] == "\n":
                            j += 1
                        buf.append("&#10;")
                    elif ch == "\n":
                        buf.append("&#10;")
                    elif ch == "\t":
                        buf.append("&#9;")
                    else:
                        buf.append(ch)
                else:
                    if ch in "\"'":
                        quote = ch
                    buf.append(ch)
                    if ch == ">":
                        j += 1
                        break
                j += 1
            out.append("".join(buf))
            i = j
    return "".join(out)


_GLUED_ATTR = re.compile(r'"(?=[A-Za-z_][\w.:-]*\s*=\s*")')


def parse_xml_file(path):
    raw = protect_attr_whitespace(open(path, "rb").read().decode("utf-8-sig"))
    try:
        return ET.fromstring(raw)
    except ET.ParseError:
        # Arcen інколи пише атрибути без пробілу: id="X"text="..." — їхній парсер це прощає, ElementTree ні.
        fixed = _GLUED_ATTR.sub('" ', raw)
        if fixed == raw:
            raise
        root = ET.fromstring(fixed)
        print(f"  (виправлено склеєні атрибути в {os.path.basename(path)})")
        return root


def truthy(v):
    return str(v).strip().lower() in ("true", "1")


# ---------------------------------------------------------------- метадані

class Schema:
    def __init__(self, name):
        self.name = name
        self.localized = {}
        self.id_attr = "id"
        self.explicit_id = False
        self.subs = {}

    def load(self, el):
        for ch in el:
            if ch.tag == "attribute":
                key = ch.get("key")
                if ch.get("is_localized") is not None:
                    self.localized[key] = truthy(ch.get("is_localized"))
                elif ch.get("type") != "existing-override":
                    self.localized.setdefault(key, False)
                if truthy(ch.get("is_id_for_layer", "false")) and not self.explicit_id:
                    self.id_attr, self.explicit_id = key, True
            elif ch.tag == "sub_node":
                sub = self.subs.get(ch.get("id")) or Schema(ch.get("id"))
                sub.load(ch)
                self.subs[ch.get("id")] = sub


class Rec:
    def __init__(self, tag, file):
        self.tag = tag
        self.file = file
        self.attrs = {}
        self.children = []

    def merge(self, el, schema, file):
        self.attrs.update(el.attrib)
        for ch in el:
            if not isinstance(ch.tag, str):
                continue
            sub = schema.subs.get(ch.tag) if schema else None
            idk = sub.id_attr if sub else "id"
            cid = ch.get(idk)
            target = None
            if cid is not None and truthy(ch.get("is_partial_record", "false")):
                target = next((c for c in self.children
                               if c.tag == ch.tag and c.attrs.get(idk) == cid), None)
            if target is None:
                target = Rec(ch.tag, file)
                self.children.append(target)
            target.merge(ch, sub, file)


# ---------------------------------------------------------------- витягування англійського

def extract_english(game):
    """-> OrderedDict key -> {"en", "file", "note"} у порядку файлів гри."""
    shared = parse_xml_file(os.path.join(game.cfg, "SharedMetaData.metadata"))
    out = collections.OrderedDict()

    def emit(prefix, rec, schema, skip, note, attr_prefix=None):
        skip = skip or truthy(rec.attrs.get("skip_all_localization_on_node", "false"))
        note = rec.attrs.get("translation_notes") or note
        if not skip:
            ap = attr_prefix or prefix
            for k, v in rec.attrs.items():
                if schema.localized.get(k) and v.strip():
                    out[ap + "|" + k] = {"en": v, "file": rec.file, "note": note or ""}
        for ch in rec.children:
            sub = schema.subs.get(ch.tag)
            if sub is not None:
                emit(f"{prefix}|{ch.tag}|{ch.attrs.get(sub.id_attr, '')}", ch, sub, skip, note)

    for folder in sorted(os.listdir(game.cfg)):
        base = os.path.join(game.cfg, folder)
        if not os.path.isdir(base):
            continue
        metas = sorted(f for f in os.listdir(base) if f.endswith(".metadata"))
        if not metas:
            continue
        meta_root = parse_xml_file(os.path.join(base, metas[0]))
        schema = Schema(meta_root.get("node_name"))
        schema.load(shared)
        schema.load(meta_root)

        dirs = [base] + [os.path.join(e, folder) for e in game.expansions
                         if os.path.isdir(os.path.join(e, folder))]
        recs, order = {}, []
        for d in dirs:
            for fn in sorted(os.listdir(d)):
                if not fn.lower().endswith(".xml"):
                    continue
                path = os.path.join(d, fn)
                rel = os.path.relpath(path, game.root)
                try:
                    root = parse_xml_file(path)
                except ET.ParseError as ex:
                    print(f"  ! не вдалося розібрати {rel}: {ex}")
                    continue
                for el in root:
                    if el.tag != schema.name or el.get("id") is None:
                        continue
                    rid = el.get("id")
                    if truthy(el.get("is_partial_record", "false")) and rid in recs:
                        recs[rid].merge(el, schema, rel)
                        continue
                    rec = Rec(el.tag, rel)
                    rec.merge(el, schema, rel)
                    if rid not in recs:
                        order.append(rid)
                    recs[rid] = rec

        table = TABLE_NAME_OVERRIDES.get(folder, re.sub(r"^\d_", "", folder))
        for rid in order:
            rec = recs[rid]
            if folder == "0_Language":
                # Таблиця мовних рядків: ключ — просто id
                if rec.attrs.get("text", "").strip() and not truthy(
                        rec.attrs.get("skip_all_localization_on_node", "false")):
                    out[rid] = {"en": rec.attrs["text"], "file": rec.file,
                                "note": rec.attrs.get("translation_notes", "")}
            elif rec.attrs.get("tl") == "OC" and rid.startswith("Cont_"):
                # Старі "споглядання": поля кореня експортуються як ContemplationType|<id без Cont_>
                emit(f"{table}|{rid}", rec, schema, False, "",
                     attr_prefix=f"ContemplationType|{rid[5:]}")
            else:
                emit(f"{table}|{rid}", rec, schema, False, "")
    return out


def load_links(game):
    """Ключі з лівого боку ALL-LOCA-LINKS.txt беруть переклад з правого — перекладати їх не треба."""
    links = {}
    path = os.path.join(game.imports, "ALL-LOCA-LINKS.txt")
    if os.path.exists(path):
        for line in open(path, encoding="utf-8-sig"):
            line = line.strip()
            if "中" in line:
                a, b = line.split("中", 1)
                links[a] = b
    return links


def load_official(game, lang, links):
    path = os.path.join(game.imports, lang + ".xlsx")
    if not os.path.exists(path):
        return {}
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    d = {}
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i and row and row[0] is not None and len(row) > 1 and row[1] is not None:
            d[str(row[0])] = str(row[1])
    wb.close()
    for a, b in links.items():
        if a not in d and b in d:
            d[a] = d[b]
    return d


# ---------------------------------------------------------------- робочий файл

def norm_text(v):
    if v is None:
        return ""
    return str(v).replace("\r\n", "\n").replace("\r", "\n")


def read_sheet(ws, key_col=COL_KEY):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    res = []
    for r in rows[1:]:
        d = {header[i]: norm_text(r[i]) for i in range(min(len(header), len(r)))}
        if d.get(key_col):
            res.append(d)
    return res


def read_work(path):
    if not os.path.exists(path):
        return [], [], {}
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    main = read_sheet(wb[SHEET_MAIN]) if SHEET_MAIN in wb.sheetnames else []
    removed = read_sheet(wb[SHEET_REMOVED]) if SHEET_REMOVED in wb.sheetnames else []
    info = {}
    if SHEET_INFO in wb.sheetnames:
        for r in wb[SHEET_INFO].iter_rows(values_only=True):
            if r and r[0]:
                info[str(r[0])] = "" if len(r) < 2 or r[1] is None else str(r[1])
    wb.close()
    return main, removed, info


def set_cell(ws, row, col, value):
    value = ILLEGAL_XLSX_RE.sub("", value) if isinstance(value, str) else value
    c = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str) and value.startswith("="):
        c.data_type = "s"  # інакше Excel вважатиме це формулою
    return c


def write_table(ws, columns, rows, widths, wrap_cols=()):
    bold = Font(bold=True)
    for ci, name in enumerate(columns, 1):
        set_cell(ws, 1, ci, name).font = bold
    wrap = Alignment(wrap_text=True, vertical="top")
    top = Alignment(vertical="top")
    for ri, r in enumerate(rows, 2):
        for ci, name in enumerate(columns, 1):
            c = set_cell(ws, ri, ci, r.get(name, ""))
            c.alignment = wrap if name in wrap_cols else top
    for ci, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(ci)].width = w
    ws.freeze_panes = "B2"
    last = openpyxl.utils.get_column_letter(len(columns))
    ws.auto_filter.ref = f"A1:{last}{max(1, len(rows) + 1)}"


def main_columns():
    return [COL_KEY, COL_EN, COL_UK, COL_STATUS, COL_PLACE] + [ref_col(l) for l in REFERENCE_LANGS] + \
           [COL_BY, COL_OLD_EN, COL_NOTE, COL_FILE]


def ref_col(lang):
    return f"{lang} (довідка)"


def save_work(path, main_rows, removed_rows, info):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_MAIN
    cols = main_columns()
    try:
        places = ui_places(main_rows)
        for r in main_rows:
            r[COL_PLACE] = places.get(r[COL_KEY], "")
    except Exception as ex:  # аналіз коду не критичний
        print(f"  ! не вдалося визначити місце в інтерфейсі: {ex}")
    widths = [38, 60, 60, 10, 9] + [45] * len(REFERENCE_LANGS) + [12, 45, 30, 30]
    write_table(ws, cols, main_rows, widths,
                wrap_cols={COL_EN, COL_UK, COL_OLD_EN, COL_NOTE} | {ref_col(l) for l in REFERENCE_LANGS})
    n = len(main_rows) + 1
    st = openpyxl.utils.get_column_letter(cols.index(COL_STATUS) + 1)
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{st}2:{st}{n}")
    for status, color in (("new", "F8D7DA"), ("changed", "FFE5B4"), ("auto", "FFF3BF"), ("mt", "DCEBFA")):
        ws.conditional_formatting.add(
            f"A2:{openpyxl.utils.get_column_letter(len(cols))}{n}",
            FormulaRule(formula=[f'${st}2="{status}"'],
                        fill=PatternFill("solid", start_color=color, end_color=color)))

    ws2 = wb.create_sheet(SHEET_REMOVED)
    write_table(ws2, [COL_KEY, COL_EN, COL_UK, "Видалено"], removed_rows, [38, 60, 60, 14],
                wrap_cols={COL_EN, COL_UK})
    ws3 = wb.create_sheet(SHEET_INFO)
    for i, (k, v) in enumerate(info.items(), 1):
        set_cell(ws3, i, 1, k)
        set_cell(ws3, i, 2, v)
    ws3.column_dimensions["A"].width = 28
    ws3.column_dimensions["B"].width = 100
    try:
        wb.save(path)
    except PermissionError:
        sys.exit(f"Не вдалося записати {path} — закрийте його в Excel і повторіть.")
    if path == WORK_FILE:
        dump_work(main_rows, removed_rows, info)


# ---------------------------------------------------------------- текстова копія для git

TEXT_DIR = os.path.join(HERE, "text")
# Похідні колонки (Місце, довідкові переклади) не зберігаємо: sync їх перераховує, а в diff вони лише шумлять.
TEXT_COLS = [COL_EN, COL_UK, COL_STATUS, COL_BY, COL_OLD_EN, COL_NOTE, COL_FILE]


def write_json(path, data):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8", newline="\n") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
        f.write("\n")
    os.replace(tmp, path)


def dump_work(main_rows, removed_rows, info):
    """text/translation.json і text/removed.json — текстова копія uk_translation.xlsx для git diff."""
    def pack(r, cols):
        return {c: r[c] for c in cols if r.get(c)}
    write_json(os.path.join(TEXT_DIR, "translation.json"),
               {r[COL_KEY]: pack(r, TEXT_COLS) for r in main_rows})
    write_json(os.path.join(TEXT_DIR, "removed.json"),
               {r[COL_KEY]: pack(r, [COL_EN, COL_UK, "Видалено"]) for r in removed_rows})
    write_json(os.path.join(TEXT_DIR, "info.json"), dict(info))


def backup_work():
    if not os.path.exists(WORK_FILE):
        return
    bdir = os.path.join(HERE, "backup")
    os.makedirs(bdir, exist_ok=True)
    stamp = datetime.datetime.now().strftime("%Y-%m-%d_%H%M%S")
    shutil.copy2(WORK_FILE, os.path.join(bdir, f"uk_translation_{stamp}.xlsx"))
    old = sorted(glob.glob(os.path.join(bdir, "uk_translation_*.xlsx")))
    for f in old[:-20]:
        os.remove(f)


# ---------------------------------------------------------------- sync

def cmd_sync(game, args):
    version = game.version()
    build = game.build_id()
    print(f"Гра: {game.root}\nВерсія: {version}, збірка Steam: {build or '?'}\nВитягую англійський текст з XML...")
    en = extract_english(game)
    links = load_links(game)
    linked = [k for k in en if k in links]
    current = collections.OrderedDict((k, v) for k, v in en.items() if k not in links)
    print(f"  рядків у грі: {len(en)}, з них через ALL-LOCA-LINKS: {len(linked)}, "
          f"до перекладу: {len(current)}")

    refs = {lang: load_official(game, lang, links) for lang in REFERENCE_LANGS}
    old_main, old_removed, _ = read_work(WORK_FILE)
    old = {r[COL_KEY]: r for r in old_main}
    removed_pool = {r[COL_KEY]: r for r in old_removed}

    # Пам'ять перекладів: однаковий англійський текст -> вже готовий переклад
    tm = {}
    for r in old_main:
        if r.get(COL_UK, "").strip() and r.get(COL_STATUS, "") in ("ok", "", "new", "mt"):
            tm.setdefault(r[COL_EN], r[COL_UK])

    new_rows, report_new, report_changed, report_restored = [], [], [], []
    for key, info in current.items():
        en_text = info["en"]
        prev = old.pop(key, None)
        restored = False
        if prev is None and key in removed_pool:
            prev, restored = removed_pool.pop(key), True
        if prev is None:
            uk = tm.get(en_text, "")
            status, old_en = ("auto" if uk else "new"), ""
            report_new.append({COL_KEY: key, COL_EN: en_text, COL_UK: uk, COL_FILE: info["file"]})
        else:
            uk = prev.get(COL_UK, "")
            status = prev.get(COL_STATUS, "").strip().lower()
            old_en = prev.get(COL_OLD_EN, "")
            if status in ("ok", ""):
                old_en = ""
            if prev.get(COL_EN, "") != en_text:
                if not uk.strip():
                    status, old_en = "new", ""
                else:
                    if status != "changed" or not old_en:
                        old_en = prev.get(COL_EN, "")
                    status = "changed"
                report_changed.append({COL_KEY: key, "Було English": prev.get(COL_EN, ""),
                                       "Стало English": en_text, COL_UK: uk,
                                       COL_FILE: info["file"]})
            if status == "changed" and old_en == en_text:
                status, old_en = "ok", ""
            if not uk.strip():
                tm_uk = tm.get(en_text, "")
                uk, status, old_en = (tm_uk, "auto", "") if tm_uk else ("", "new", "")
            elif status not in STATUSES:
                status = "ok"
            elif status == "new":
                status = "ok"
            if restored:
                report_restored.append({COL_KEY: key, COL_EN: en_text, COL_UK: uk})
        row = {COL_KEY: key, COL_EN: en_text, COL_UK: uk, COL_STATUS: status,
               COL_BY: (prev or {}).get(COL_BY, "") if uk else "",
               COL_OLD_EN: old_en, COL_NOTE: info["note"], COL_FILE: info["file"]}
        for lang in REFERENCE_LANGS:
            row[ref_col(lang)] = refs[lang].get(key, "")
        new_rows.append(row)

    today = datetime.date.today().isoformat()
    report_removed = []
    for key, r in old.items():
        report_removed.append({COL_KEY: key, COL_EN: r.get(COL_EN, ""), COL_UK: r.get(COL_UK, "")})
        removed_pool[key] = {COL_KEY: key, COL_EN: r.get(COL_EN, ""), COL_UK: r.get(COL_UK, ""),
                             "Видалено": today}
    # видалені без перекладу не тримаємо
    removed_rows = [r for r in removed_pool.values() if r.get(COL_UK, "").strip()]

    first_run = not old_main
    print(f"  нових: {len(report_new)}, змінених: {len(report_changed)}, "
          f"видалених: {len(report_removed)}, повернених: {len(report_restored)}")
    if args.dry_run:
        print("(--dry-run: нічого не записано)")
        return

    backup_work()
    info = collections.OrderedDict([
        ("Версія гри", version), ("Збірка Steam", build or ""),
        ("Синхронізовано", datetime.datetime.now().isoformat(" ", "seconds")),
        ("Тека гри", game.root), ("Рядків до перекладу", str(len(current))),
        ("Рядків через ALL-LOCA-LINKS", str(len(linked))),
        ("Статуси", "new — не перекладено; changed — англійський текст змінився, перевірте переклад "
                    "(старий текст у колонці «Попередній English»), потім поставте ok; "
                    "auto — підставлено з такого самого англійського рядка, перевірте і поставте ok; "
                    "mt — машинний переклад (ШІ), потребує редактури; "
                    "ok — готово."),
    ])
    save_work(WORK_FILE, new_rows, removed_rows, info)
    print(f"Робочий файл: {WORK_FILE}")

    snap_dir = os.path.join(HERE, "snapshots")
    os.makedirs(snap_dir, exist_ok=True)
    # id версії гри не змінюється з кожним патчем, тож знімок іменуємо ще й збіркою Steam,
    # щоб нова збірка не перезаписала знімок попередньої
    snap = os.path.join(snap_dir, f"en_{version}" + (f"_b{build}" if build else "") + ".json")
    with open(snap, "w", encoding="utf-8", newline="\n") as f:
        json.dump({k: v["en"] for k, v in en.items()}, f, ensure_ascii=False, indent=0)

    if not first_run and (report_new or report_changed or report_removed or report_restored):
        ch_dir = os.path.join(HERE, "changes")
        os.makedirs(ch_dir, exist_ok=True)
        stamp = datetime.datetime.now().strftime("%Y-%m-%d_%H%M%S")
        path = os.path.join(ch_dir, f"{stamp}_{version}.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Нові"
        write_table(ws, [COL_KEY, COL_EN, COL_UK, COL_FILE], report_new, [38, 70, 50, 30],
                    wrap_cols={COL_EN, COL_UK})
        write_table(wb.create_sheet("Змінені"),
                    [COL_KEY, "Було English", "Стало English", COL_UK, COL_FILE], report_changed,
                    [38, 60, 60, 50, 30], wrap_cols={"Було English", "Стало English", COL_UK})
        write_table(wb.create_sheet("Видалені"), [COL_KEY, COL_EN, COL_UK], report_removed,
                    [38, 70, 50], wrap_cols={COL_EN, COL_UK})
        if report_restored:
            write_table(wb.create_sheet("Повернені"), [COL_KEY, COL_EN, COL_UK], report_restored,
                        [38, 70, 50], wrap_cols={COL_EN, COL_UK})
        wb.save(path)
        print(f"Звіт змін: {path}")
    print_status(new_rows)


# ---------------------------------------------------------------- build

# ---------------------------------------------------------------- де текст показується в інтерфейсі

PLACE_BUTTON, PLACE_NAME, PLACE_TEXT = "кнопка", "назва", "текст"
COL_PLACE = "Місце"
TIGHT_FIELDS = {"short_name", "very_short_text", "name_prefix", "equipment_prefix", "toast_line_1",
                "toast_line_2", "progress_label_text", "popup_text_on_complete", "popup_text_on_disband",
                "alternate_maximum_label_for_tooltip"}
TEXT_FIELDS = {"description", "strategy_tip", "strategy_tip_optional", "result_message", "how_to_acquire",
               "technical_notes", "short_description", "result_message_secondary_note", "high_text",
               "alt_description_on_map_tooltip", "alt_strategy_tip_on_map_tooltip", "benefit_reason",
               "extra_lore_info_optional", "method_details", "alt_description", "alt_strategy_tip",
               "toast_tooltip_instruction", "line_2", "restrictions", "trigger_from_event_note",
               "undiscovered_body", "network_description_optional"}
TIGHT_TABLES = {"InputCode", "SettingCategory", "SettingSubcategory", "MachineBuildModeCategory",
                "MachineCommandModeCategory", "MachineVRModeCategory", "UISidebarType", "UISidebarBasicItem",
                "InputActionCategory", "EquipmentSlotType", "CityLensType"}
_UI_CODE_KINDS = None


def ui_code_kinds(game_root):
    """Статичний аналіз C#-коду інтерфейсу: як показується кожен мовний ключ."""
    global _UI_CODE_KINDS
    if _UI_CODE_KINDS is not None:
        return _UI_CODE_KINDS
    sig_re = re.compile(r"^\s*(?:(?:public|private|protected|internal|static|override|virtual|sealed|async|"
                        r"unsafe|new)\s+)+[\w<>\[\],\.\s]+?\s+(\w+)\s*\(")
    call_re = re.compile(r"(\w*Lang\w*)\(\s*\"([^\"]+)\"")
    rank = {PLACE_BUTTON: 3, PLACE_NAME: 2, "?": 1, PLACE_TEXT: 0}
    kinds = {}
    files = glob.glob(os.path.join(game_root, "CodeExternal", "**", "*.cs"), recursive=True) + \
        glob.glob(os.path.join(game_root, "CodeMain", "**", "*.cs"), recursive=True)
    for f in files:
        depth, method, method_depth, pending = 0, None, None, None
        for line in open(f, encoding="utf-8-sig", errors="replace"):
            m = sig_re.match(line)
            if m and not line.strip().endswith(";"):
                pending = m.group(1)
            for ch in line:
                if ch == "{":
                    depth += 1
                    if pending:
                        method, method_depth, pending = pending, depth, None
                elif ch == "}":
                    if method and depth == method_depth:
                        method = None
                    depth -= 1
            for cm in call_re.finditer(line):
                if re.search(r"novel\.(TitleUpperLeft|TitleLowerLeft|TitleUpperRight|FrameTitle|MainHeader)\b", line):
                    k = PLACE_NAME
                elif "AfterLineItemHeader" in cm.group(1):
                    k = PLACE_NAME
                elif re.search(r"novel\.(Main|FrameBody|StatWorking\w*)\b", line):
                    k = PLACE_TEXT
                elif method and re.search(r"FillTextPerFrame|FillOtherTextPerFrame|GetTextToShow|SetText|FillText",
                                          method):
                    k = PLACE_BUTTON
                else:
                    k = "?"
                old = kinds.get(cm.group(2))
                if old is None or rank[k] > rank[old]:
                    kinds[cm.group(2)] = k
    _UI_CODE_KINDS = kinds
    return kinds


def ui_place(key, en, code_kinds):
    """кнопка — тісне місце (кнопка, вкладка, мітка); назва — заголовок/назва в списку; текст — абзаци."""
    v = visible_text(en)
    short = "\n" not in v and len(v) <= 30
    parts = key.split("|")
    if len(parts) == 1:  # мовний рядок
        k = code_kinds.get(key)
        if k == PLACE_BUTTON:
            return PLACE_BUTTON if short else PLACE_TEXT
        if k in (PLACE_NAME, PLACE_TEXT):
            return k if short or k == PLACE_TEXT else PLACE_TEXT
        if short and re.search(r"Abbrev|_Short|Brief|Button|_Tab|Header|Label|Title", key):
            return PLACE_BUTTON
        if short and not re.search(r"[.!?:]\s*$", v):
            return PLACE_NAME
        return PLACE_TEXT
    table, field = parts[0], parts[-1]
    if field in TIGHT_FIELDS or "abbrev" in field:
        return PLACE_BUTTON
    if field in TEXT_FIELDS or not short:
        return PLACE_TEXT
    if field == "display_name" and len(parts) == 3 and (
            table in TIGHT_TABLES or table.endswith("Collection") or table.startswith("UIX_")):
        return PLACE_BUTTON
    return PLACE_NAME


def ui_places(rows):
    kinds = ui_code_kinds(GAME_ROOT or find_game_dir(None))
    return {r[COL_KEY]: ui_place(r[COL_KEY], r.get(COL_EN, ""), kinds) for r in rows}


def too_long(place, en, uk):
    e, u = len(visible_text(en)), len(visible_text(uk))
    if place == PLACE_BUTTON:
        return e > 0 and u > max(e + 3, e * 1.25 + 1)
    if place == PLACE_NAME:
        return e > 0 and u > e * 1.6 + 3
    return False


def visible_text(s):
    """Текст без тегів; плейсхолдер рахується як один символ."""
    return PLACEHOLDER_RE.sub("0", TAG_RE.sub("", s)).strip()


_FORBIDDEN = None


_SUSPECT = None


def _load_rules(path):
    errors, suspects = [], []
    if os.path.exists(path):
        for line in open(path, encoding="utf-8-sig"):
            line = line.rstrip("\r\n")
            if not line.strip() or line.lstrip().startswith("#"):
                continue
            suspect = line.startswith("?")
            pat, _, why = line.lstrip("?").partition("\t")
            try:
                rx = re.compile(r"(?<![^\W_])(?:" + pat.strip() + r")(?![^\W_])", re.I)
            except re.error as ex:
                print(f"  ! поганий вираз у {os.path.basename(path)}: {pat.strip()} ({ex})")
                continue
            (suspects if suspect else errors).append((rx, why.strip() or pat.strip()))
    return errors, suspects


def forbidden_rules():
    """Помилки: forbidden_uk.txt і рядки без «?» у russisms_uk.txt."""
    global _FORBIDDEN, _SUSPECT
    if _FORBIDDEN is None:
        e1, s1 = _load_rules(os.path.join(HERE, "forbidden_uk.txt"))
        e2, s2 = _load_rules(os.path.join(HERE, "russisms_uk.txt"))
        _FORBIDDEN, _SUSPECT = e1 + e2, s1 + s2
    return _FORBIDDEN


def suspect_rules():
    """Підозри («?» у russisms_uk.txt): лише попередження для редактора."""
    forbidden_rules()
    return _SUSPECT


def suspect_issues(uk):
    out = []
    for rx, why in suspect_rules():
        m = rx.search(uk)
        if m:
            out.append(f"«{m.group(0)}» — {why}")
    return out


def check_row(key, en, uk):
    issues = []
    for rx, why in forbidden_rules():
        m = rx.search(uk)
        if m:
            issues.append(f"заборонене слово «{m.group(0)}»: {why}")
    if sorted(PLACEHOLDER_RE.findall(en)) != sorted(PLACEHOLDER_RE.findall(uk)):
        issues.append("плейсхолдери {..} не збігаються")
    if sorted(TAG_RE.findall(en)) != sorted(TAG_RE.findall(uk)):
        issues.append("теги <..> не збігаються")
    if en[:1].isspace() != uk[:1].isspace() or en[-1:].isspace() != uk[-1:].isspace():
        issues.append("пробіл/перенос на початку чи в кінці відрізняється")
    if en.count("\n") != uk.count("\n"):
        issues.append(f"кількість рядків: en {en.count(chr(10)) + 1}, uk {uk.count(chr(10)) + 1}")
    return issues


def xml_text(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def write_arcen_xlsx(game, pairs, path):
    """Пише xlsx у тому самому мінімальному форматі, що й офіційні файли гри
    (гра читає їх власним парсером: xl/sharedStrings.xml + xl/worksheets/sheet1.xml)."""
    template = next((os.path.join(game.imports, l + ".xlsx") for l in OFFICIAL_LANGS
                     if os.path.exists(os.path.join(game.imports, l + ".xlsx"))), None)
    if template is None:
        sys.exit("Не знайдено жодного офіційного xlsx як шаблону.")
    strings = ["Key", "Text"]
    for k, v in pairs:
        strings += [k, v]
    sst = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
           '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
           f'count="{len(strings)}" uniqueCount="{len(strings)}">']
    for s in strings:
        s = s.replace("\r\n", "\n").replace("\r", "\n")
        t = xml_text(s).replace("\n", "\r\n")
        pre = ' xml:space="preserve"' if s != s.strip() else ""
        sst.append(f"<si><t{pre}>{t}</t></si>")
    sst.append("</sst>")

    nrows = len(pairs) + 1
    sheet = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
             '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
             'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
             f'<dimension ref="A1:B{nrows}"/><sheetViews><sheetView tabSelected="1" workbookViewId="0">'
             '<pane ySplit="1" topLeftCell="A2" activePane="bottomLeft" state="frozen"/>'
             '<selection pane="bottomLeft" activeCell="A2" sqref="A2"/></sheetView></sheetViews>'
             '<sheetFormatPr defaultRowHeight="14.3"/><cols>'
             '<col min="1" max="1" width="58.125" customWidth="1"/>'
             '<col min="2" max="2" width="86.375" customWidth="1" style="2"/></cols><sheetData>'
             '<row r="1"><c r="A1" s="1" t="s"><v>0</v></c><c r="B1" s="3" t="s"><v>1</v></c></row>']
    for i in range(len(pairs)):
        r = i + 2
        sheet.append(f'<row r="{r}"><c r="A{r}" t="s"><v>{2 * i + 2}</v></c>'
                     f'<c r="B{r}" t="s"><v>{2 * i + 3}</v></c></row>')
    sheet.append('</sheetData><pageMargins left="0.7" right="0.7" top="0.75" bottom="0.75" '
                 'header="0.3" footer="0.3"/></worksheet>')

    tmp = path + ".tmp"
    with zipfile.ZipFile(template) as zin, \
            zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            if item.filename == "xl/sharedStrings.xml":
                zout.writestr(item.filename, "".join(sst).encode("utf-8"))
            elif item.filename == "xl/worksheets/sheet1.xml":
                zout.writestr(item.filename, "".join(sheet).encode("utf-8"))
            else:
                zout.writestr(item.filename, zin.read(item.filename))
    try:
        os.replace(tmp, path)
    except PermissionError:
        os.remove(tmp)
        sys.exit(f"Не вдалося записати {path} — можливо, файл відкритий або гра запущена.")

    # перевірка: читаємо назад
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    back = [(str(r[0]), norm_text(r[1])) for i, r in
            enumerate(wb[wb.sheetnames[0]].iter_rows(values_only=True)) if i]
    wb.close()
    expected = [(k, v.replace("\r\n", "\n").replace("\r", "\n")) for k, v in pairs]
    if back != expected:
        sys.exit("Помилка: записаний xlsx не збігається з даними при зворотному читанні!")


def ensure_localization_type(game):
    if not os.path.exists(game.loctypes):
        print(f"  ! не знайдено {game.loctypes}")
        return
    text = open(game.loctypes, "rb").read().decode("utf-8-sig")
    nl = "\r\n" if "\r\n" in text else "\n"
    block = open(LOCTYPE_BLOCK_FILE, encoding="utf-8").read().strip("\r\n")
    block = block.replace("\r\n", "\n").replace("\n", nl)
    existing = re.search(r'[ \t]*<localization_type\s+id="%s".*?/>' % LANG_ID, text, re.S)
    if existing and existing.group(0).strip() == block.strip():
        print("LocalizationTypes.xml: українська вже є.")
        return
    bdir = os.path.join(HERE, "backup")
    os.makedirs(bdir, exist_ok=True)
    shutil.copy2(game.loctypes, os.path.join(bdir, f"LocalizationTypes_{game.version()}.xml"))
    if existing:
        # правки в localization_type_uk.xml (колонки, роздільники тощо) переносимо в гру
        text = text[:existing.start()] + block + text[existing.end():]
        msg = "LocalizationTypes.xml: запис української оновлено з localization_type_uk.xml."
    else:
        idx = text.rfind("</root>")
        if idx < 0:
            print("  ! LocalizationTypes.xml має неочікуваний формат, пропускаю")
            return
        text = text[:idx].rstrip() + nl + nl + block + nl + nl + text[idx:]
        msg = "LocalizationTypes.xml: українську мову додано (оновлення гри її стерло)."
    with open(game.loctypes, "w", encoding="utf-8", newline="") as f:
        f.write(text)
    print(msg)


def link_duplicates(game, rows, pairs, verbose=False):
    """ALL-LOCA-LINKS посилається на NPCEvent|Cont_X|поле, а старі споглядання ми експортуємо
    як ContemplationType|X|поле — дублюємо переклад під ключем цілі (зайві ключі гра ігнорує)."""
    keys, done = {r[COL_KEY] for r in rows}, dict(pairs)
    targets = sorted({b for b in load_links(game).values()
                      if b.startswith("NPCEvent|Cont_") and b not in keys})
    src = {t: "ContemplationType|" + t[len("NPCEvent|Cont_"):] for t in targets}
    extra = [(t, done[src[t]]) for t in targets if src[t] in done]
    if verbose:
        print(f"Цілей ALL-LOCA-LINKS виду NPCEvent|Cont_… без власного рядка: {len(targets)}, "
              f"з перекладом з ContemplationType: {len(extra)}")
        for t in targets:
            if src[t] not in done:
                print(f"  без перекладу: {t}")
    return extra


def cmd_build(game, args):
    rows, removed, info = read_work(WORK_FILE)
    if not rows:
        sys.exit(f"Немає {WORK_FILE}. Спершу запустіть:  python hotm_uk.py sync")
    dump_work(rows, removed, info)  # ручні правки в Excel теж потрапляють у text/
    pairs, warnings = [], []
    skipped_changed = 0
    counts = collections.Counter()
    for r in rows:
        uk = r.get(COL_UK, "")
        status = r.get(COL_STATUS, "")
        counts[status] += 1
        if not uk.strip():
            continue
        if status == "changed" and args.skip_changed:
            skipped_changed += 1
            continue
        for issue in check_row(r[COL_KEY], r.get(COL_EN, ""), uk):
            warnings.append((r[COL_KEY], issue, r.get(COL_EN, ""), uk))
        pairs.append((r[COL_KEY], uk))

    extra = link_duplicates(game, rows, pairs, verbose=True)
    write_arcen_xlsx(game, pairs + extra, game.out_xlsx)
    print(f"Записано {game.out_xlsx}: {len(pairs)} рядків з {len(rows)} "
          f"({100 * len(pairs) / max(1, len(rows)):.1f}%) + {len(extra)} дублікатів для ALL-LOCA-LINKS")
    if counts["changed"]:
        print(f"  увага: {counts['changed']} рядків зі статусом changed "
              + ("не включено (--skip-changed)" if args.skip_changed else "включено зі старим перекладом"))
    if counts["auto"]:
        print(f"  {counts['auto']} рядків зі статусом auto (підставлено автоматично) — варто перевірити")
    spath = os.path.join(HERE, "build_suspects.txt")
    n_susp = 0
    with open(spath, "w", encoding="utf-8") as f:
        f.write("Можливі русизми й кальки (russisms_uk.txt, рядки з «?»). Вирішує редактор.\n\n")
        for key, uk in pairs:
            for si in suspect_issues(uk):
                n_susp += 1
                f.write(f"{key}\n  {si}\n  UK: {uk[:200]!r}\n\n")
    if n_susp:
        print(f"  можливих русизмів/кальок: {n_susp} — див. {spath}")
    wpath = os.path.join(HERE, "build_warnings.txt")
    with open(wpath, "w", encoding="utf-8") as f:
        for key, issue, en, uk in warnings:
            f.write(f"{key}\n  {issue}\n  EN: {en!r}\n  UK: {uk!r}\n\n")
    if warnings:
        print(f"  попереджень перевірки: {len(warnings)} — див. {wpath}")

    # Кнопки/мітки й назви, помітно довші за англійські (кирилиця йде ширшим запасним шрифтом)
    by_key = {r[COL_KEY]: r for r in rows}
    places = ui_places(rows)
    long_ui = []
    for key, uk in pairs:
        en = by_key[key].get(COL_EN, "")
        place = places.get(key, PLACE_TEXT)
        if too_long(place, en, uk):
            e, u = visible_text(en), visible_text(uk)
            long_ui.append((place != PLACE_BUTTON, -len(u) / max(1, len(e)), key, place, e, u))
    long_ui.sort()
    lpath = os.path.join(HERE, "build_long_ui.txt")
    with open(lpath, "w", encoding="utf-8") as f:
        f.write("Написи, які можуть не вміститися (кирилиця — ширший шрифт). Спершу кнопки/мітки, потім назви.\n"
                "Описи й підказки (текст) тут не перевіряються — вони можуть бути довгими.\n\n")
        for _, neg, key, place, e, u in long_ui:
            f.write(f"{-neg:4.1f}x [{place}] {key}\n       EN ({len(e)}): {e}\n       UK ({len(u)}): {u}\n")
    if long_ui:
        print(f"  можливо задовгих коротких написів: {len(long_ui)} — див. {lpath}")
    ensure_localization_type(game)


def cmd_pseudo(game, args):
    en = extract_english(game)
    links = load_links(game)
    pairs = [(k, ("ЇҐ·" + v["en"]) if re.search(r"[A-Za-z]{2}", v["en"]) else v["en"])
             for k, v in en.items() if k not in links]
    write_arcen_xlsx(game, pairs, game.out_xlsx)
    print(f"Тестовий {game.out_xlsx}: {len(pairs)} рядків з префіксом «ЇҐ·».")
    print("Після перевірки зберіть справжній файл командою build.")
    ensure_localization_type(game)


# ---------------------------------------------------------------- глосарій

GLOSSARY_FILE = os.path.join(HERE, "glossary.xlsx")
G_EN, G_UK, G_VAR, G_NOTE, G_STATUS = "English", "Українська", "Варіанти", "Примітка", "Статус"
G_FREQ, G_CAT, G_DE, G_EX = "Частота", "Категорія", "German (довідка)", "Приклад"
G_COLS = [G_EN, G_UK, G_VAR, G_NOTE, G_STATUS, G_FREQ, G_CAT, G_DE, G_EX]
G_STATUSES = ("candidate", "proposed", "ok", "skip")
PROPER_NOUN = "Власна назва"

# Таблиці, чиї display_name — ігрові терміни (ресурси, фракції, юніти, стати...)
TERM_TABLES = {
    "ResourceType", "NPCCohort", "NPCCohortGroup", "NPCCohortCategory", "MachineUnitType",
    "MachineVehicleType", "MachineStructureType", "NPCUnitType", "ActorCollection", "ActorDataType",
    "ActorStatus", "ActorFeat", "ActorPerk", "ActorBadge", "AbilityType", "UnitEquipmentType",
    "EquipmentSlotType", "ResourceConsumable", "MachineJob", "MachineJobClass", "POIType",
    "DistrictType", "BuildingType", "BuildingTag", "ResearchDomain", "Unlock", "UnlockInspirationType",
    "MetaResourceType", "HackingDaemonType", "HackingAction", "BeaconType", "KeyContact", "NamedThing",
    "ExplorationSiteType", "Swarm", "SecurityClearance", "MetaChapter", "TimelineGoal",
    "CityTimelineDoomType", "CityLensType", "NPCMission", "NPCDeal", "DaringPurchase",
    "MachineCommandModeCategory", "MachineBuildModeCategory", "MachineVRModeCategory",
    "MachineVRModeAction", "MachineCommandModeAction", "NPCDebateAction", "ComplexityType",
    "AnalysisType", "LocationActionType", "ActionOverTimeType", "UpgradeInt", "UpgradeFloat",
    "CityTimelineCrossover", "NPCUnitObjective", "MachineUnitStance", "ProfessionType",
    "EconomicClassType", "TerritoryControlType", "ForcesCollection", "UnitEquipmentCollection",
}
CORE_TERM_TABLES = {
    "ResourceType", "NPCCohort", "NPCCohortGroup", "ActorDataType", "MachineUnitType",
    "MachineVehicleType", "MetaResourceType", "ResearchDomain", "SecurityClearance", "POIType",
    "DistrictType", "MachineJobClass", "EquipmentSlotType", "KeyContact", "NamedThing", "Swarm",
    "MetaChapter", "HackingDaemonType",
}


def read_glossary():
    if not os.path.exists(GLOSSARY_FILE):
        return []
    wb = openpyxl.load_workbook(GLOSSARY_FILE, read_only=True, data_only=True)
    rows = read_sheet(wb[wb.sheetnames[0]], key_col=G_EN)
    wb.close()
    return rows


def save_glossary(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Глосарій"
    write_table(ws, G_COLS, rows, [28, 30, 25, 40, 11, 9, 24, 28, 70],
                wrap_cols={G_UK, G_VAR, G_NOTE, G_EX, G_DE})
    n = len(rows) + 1
    dv = DataValidation(type="list", formula1='"' + ",".join(G_STATUSES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"E2:E{n}")
    for status, color in (("candidate", "EEEEEE"), ("proposed", "FFF3BF"), ("skip", "D9D9D9")):
        ws.conditional_formatting.add(
            f"A2:I{n}", FormulaRule(formula=[f'$E2="{status}"'],
                                    fill=PatternFill("solid", start_color=color, end_color=color)))
    try:
        wb.save(GLOSSARY_FILE)
    except PermissionError:
        sys.exit(f"Не вдалося записати {GLOSSARY_FILE} — закрийте його в Excel.")
    dump_glossary(rows)


def dump_glossary(rows):
    """text/glossary.json — текстова копія glossary.xlsx для git diff."""
    write_json(os.path.join(TEXT_DIR, "glossary.json"),
               [{c: r[c] for c in G_COLS if r.get(c)} for r in rows])


def term_candidates(game, en, links):
    texts = [v["en"] for k, v in en.items() if k not in links]
    blob = "\n".join(texts).lower()
    de = load_official(game, "German", links)
    found = collections.OrderedDict()
    for k, v in en.items():
        p = k.split("|")
        if len(p) != 3 or p[2] != "display_name" or p[0] not in TERM_TABLES:
            continue
        t = v["en"].strip()
        if not (2 < len(t) <= 40) or len(t.split()) > 5 or re.search(r"[{}<>\n]", t):
            continue
        e = found.setdefault(t.lower(), {"en": t, "tables": set(), "keys": [], "own": 0})
        e["tables"].add(p[0])
        e["keys"].append(k)
        e["own"] += 1

    def count(term):
        return len(re.findall(r"(?<![a-z0-9])" + re.escape(term.lower()) + r"(?:s|es)?(?![a-z0-9])", blob))

    res = []
    for e in found.values():
        freq = max(0, count(e["en"]) - e["own"])
        if freq >= 2 or e["tables"] & CORE_TERM_TABLES:
            res.append({G_EN: e["en"], G_FREQ: freq, G_CAT: ", ".join(sorted(e["tables"])),
                        G_DE: de.get(e["keys"][0], "")})
    # власні назви: слова з великої літери посеред речення, які майже не трапляються з малої
    cap = collections.Counter()
    for tx in texts:
        for m in re.finditer(r"(?<=[a-z,;:] )([A-Z][A-Za-z\-]{2,})", tx):
            cap[m.group(1)] += 1
    lower = collections.Counter(re.findall(r"(?<![A-Za-z])[a-z][a-z\-]+", "\n".join(texts)))
    for w, n in cap.items():
        if n >= 8 and w.lower() not in found and lower[w.lower()] < max(2, n * 0.1):
            res.append({G_EN: w, G_FREQ: n, G_CAT: PROPER_NOUN, G_DE: ""})
    for r in res:
        low = r[G_EN].lower()
        ex = next((t for t in texts if low in t.lower() and t.strip().lower() != low), "")
        r[G_EX] = ex[:220]
    return res


def cmd_glossary(game, args):
    en = extract_english(game)
    links = load_links(game)
    cands = term_candidates(game, en, links)
    old = {r[G_EN].lower(): r for r in read_glossary()}
    added = 0
    rows = []
    for c in cands:
        r = old.pop(c[G_EN].lower(), None)
        if r is None:
            r = {G_EN: c[G_EN], G_UK: "", G_VAR: "", G_NOTE: "", G_STATUS: "candidate"}
            added += 1
        r.update({G_FREQ: c[G_FREQ], G_CAT: c[G_CAT], G_DE: c[G_DE] or r.get(G_DE, ""), G_EX: c[G_EX]})
        rows.append(r)
    rows += list(old.values())  # додані вручну терміни лишаються
    for r in rows:
        try:
            r[G_FREQ] = int(r.get(G_FREQ) or 0)
        except ValueError:
            r[G_FREQ] = 0
    rows.sort(key=lambda r: -r[G_FREQ])
    save_glossary(rows)
    c = collections.Counter(r.get(G_STATUS, "") for r in rows)
    print(f"Глосарій: {GLOSSARY_FILE}\n  термінів: {len(rows)}, нових кандидатів: {added}; "
          + ", ".join(f"{k or '—'} {v}" for k, v in c.items()))


def uk_norm(s):
    s = s.lower().replace("ʼ", "'").replace("’", "'").replace("ё", "е")
    # чергування голосних в основі (бій/бою, ніч/ночі, очки/очок): і → о
    return s.replace("і", "о")


def uk_stems(term):
    """Основи слів терміна: відкидаємо закінчення, короткі слова лишаємо (2 літери) — шукаємо на початку слова."""
    stems = []
    for w in re.findall(r"[^\W_]+", uk_norm(term)):
        if len(w) <= 2:
            stems.append(w)
            continue
        cut = 3 if len(w) >= 7 else 2 if len(w) >= 5 else 1
        stems.append(w[:max(2, len(w) - cut)])
    return stems


def term_chunks(t):
    return re.findall(r"[A-Z]+(?![a-z])|[A-Z]?[a-z]+|\d+", t)


def term_norm(t):
    return re.sub(r"[^a-z0-9]", "", t.lower())


class TermMatcher:
    """Шукає терміни глосарію в англійському тексті: без урахування регістру, з множиною -s/-es,
    і нечутливо до пробілів/дефісів усередині (NanoSeed = Nano Seed = Nano-Seed)."""

    def __init__(self, entries):
        self.map = {}
        pats = {}
        for e in entries:
            chunks = term_chunks(e[G_EN])
            if not chunks:
                continue
            key = term_norm(e[G_EN])
            if key in self.map:
                continue
            self.map[key] = e
            pats[key] = r"[\s\-'’.]{0,2}".join(re.escape(c) for c in chunks)
        order = sorted(pats, key=len, reverse=True)
        self.rx = None
        if order:
            self.rx = re.compile(r"(?<![A-Za-z0-9])(" + "|".join(pats[k] + r"(?:s|es)?" for k in order)
                                 + r")(?![A-Za-z0-9])", re.I)

    def get(self, term):
        return self.map.get(term_norm(term))

    def find(self, text):
        out = []
        if not self.rx:
            return out
        for m in self.rx.finditer(text):
            w = term_norm(m.group(1))
            e = self.map.get(w)
            if e is None and w.endswith("es"):
                e = self.map.get(w[:-2])
            if e is None and w.endswith("s"):
                e = self.map.get(w[:-1])
            if e is None:
                continue
            # власні назви — лише з великої літери
            if e.get(G_CAT) == PROPER_NOUN and not m.group(1)[:1].isupper():
                continue
            if e not in out:
                out.append(e)
        return out


def term_ok(entry, uk_text):
    text = uk_norm(uk_text)
    variants = [entry.get(G_UK, "")] + re.split(r"[;/]", entry.get(G_VAR, "") or "")
    variants += re.split(r"/", entry.get(G_UK, ""))
    for v in variants:
        st = uk_stems(v)
        if st and all(re.search(r"(?<![^\W_])" + re.escape(x), text) for x in st):
            return True
    return False


def glossary_entries():
    return [r for r in read_glossary()
            if r.get(G_UK, "").strip() and r.get(G_STATUS, "") != "skip"]


def glossary_hint(e):
    note = (e.get(G_NOTE) or "").strip()
    return e[G_UK].strip() + (f" — {note}" if note else "")


# ---------------------------------------------------------------- батчі для ШІ

BATCH_ROOT = os.path.join(HERE, "batches")
STYLE_FILE = os.path.join(HERE, "STYLE_uk.md")

STRINGS_INSTRUCTIONS = (
    "Translate every item's 'en' from English into Ukrainian for the video game Heart of the Machine, "
    "following the style guide (STYLE_uk.md). Keep {placeholders}, <tags>, the number of line breaks and "
    "leading/trailing spaces exactly. Use the 'glossary' translations for terms (inflect them as needed). "
    "'k' is the data key (context: table|record|field), 'de' is the official German translation "
    "(use it as a hint for meaning and grammatical gender), 'note' is a developer note. "
    "If 'old_uk' is present, update that existing translation to match the new 'en' (old text in 'old_en'). "
    "Return ONLY one JSON object mapping every item id to its Ukrainian text: {\"1\": \"...\", \"2\": \"...\"}.")
GLOSSARY_INSTRUCTIONS = (
    "Propose Ukrainian translations for these game terms of Heart of the Machine (cyberpunk/dystopian "
    "strategy; you play an emergent machine intelligence). 'cat' = where the term comes from, 'de' = "
    "official German translation, 'ex' = usage example. Follow the style guide (STYLE_uk.md): concise, "
    "natural Ukrainian; transliterate brand/corporation/proper names; nominative case; lower case unless "
    "a proper name. Return ONLY JSON: {\"<id>\": {\"uk\": \"term\", \"note\": \"gender, plural, "
    "declension hints or why\", \"var\": \"optional accepted alternatives separated by ;\"}}. "
    "If an entry is just a generic word that needs no fixed translation, return {\"skip\": true}.")
GLOSSARY_REVIEW_INSTRUCTIONS = (
    "You are the terminology editor of the Ukrainian translation of Heart of the Machine. 'items' is the "
    "whole draft glossary (en, uk, note, var, cat = source table, freq = how often the English term occurs in "
    "game texts). It was drafted in several parts, so fix inconsistencies: the same English word or base term "
    "must be rendered the same way everywhere (singular/plural pairs, compounds that contain a base term, "
    "corporation divisions, faction names). Glossary entries are matched case-insensitively in ALL game texts and "
    "shown to translators as mandatory, so mark as {\"skip\": true} every entry whose translation would be wrong "
    "or misleading for ordinary uses of the English word (generic words translated in a narrow sense, e.g. a "
    "building tag 'Fire' = fire station). Keep real game concepts. Return ONLY JSON with entries you change: "
    "{\"<id>\": {\"uk\": \"...\", \"note\": \"...\", \"var\": \"...\"}} or {\"<id>\": {\"skip\": true}}; "
    "omit unchanged entries.")
SHORTEN_INSTRUCTIONS = (
    "You are the UI editor of the Ukrainian translation of Heart of the Machine. The game renders Cyrillic with "
    "a WIDER fallback font, so short UI labels overflow buttons, tabs and tooltips. For every item, 'uk' is the "
    "current translation of the short English label 'en'; rewrite it so its visible length (without tags and "
    "placeholders) is at most 'max_len' characters, keeping the meaning and the glossary terms. Use natural "
    "concise Ukrainian: shorter synonyms, dropping redundant words, standard abbreviations (макс., мін., од., "
    "к-сть, авт.), arrows for directions (↑ ↓ ← →) only for key names. If the current text is already the best "
    "short form or cannot be shortened without losing meaning, return it unchanged. Keep {placeholders} and "
    "<tags>. Return ONLY JSON {\"<id>\": \"Ukrainian text\"} for every item.")
FIX_INSTRUCTIONS = (
    "You are the technical editor of the Ukrainian translation of Heart of the Machine. For every item, 'uk' is "
    "the current translation of 'en' and 'issues' lists what is wrong (line-break count, tags, placeholders, "
    "leading/trailing spaces, forbidden words). Fix exactly these issues so the structure matches 'en' "
    "(same number of \\n line breaks in the same places, same tags and placeholders, same edge spaces) while "
    "keeping the wording. Return ONLY JSON {\"<id>\": \"corrected Ukrainian text\"} for every item.")
STYLE_INSTRUCTIONS = (
    "You are the literary editor of the Ukrainian translation of Heart of the Machine. For every item, 'uk' is "
    "the current Ukrainian translation of 'en'. Rewrite ONLY the strings that do not read like natural Ukrainian "
    "written by a native speaker: russisms and calques from Russian or English, bureaucratic or overly ornate "
    "phrasing, literal word order, stacked particles (оця, саме, водночас, іще), verbal nouns where a verb is "
    "natural, overlong sentences, literal renderings of English idioms, odd collocations. The result must be "
    "plain, lively, idiomatic Ukrainian with the SAME meaning and the same register as the original (noir "
    "narration, dry humour, distinct character voices, UI brevity). Keep glossary terms, {placeholders}, <tags>, "
    "the number of line breaks and leading/trailing spaces; address the player with «ви»; game points are "
    "«одиниці», never «очки». Also check every sentence against the syntax checklist in STYLE_uk.md "
    "(section «Синтаксис: 12 хиб»): needless «щоб» before an infinitive or where English 'to' means a result; "
    "chains of subordinate clauses; calqued questions; filler pronouns («те, що», «той факт, що») and "
    "possessives («свій», «ваш» with body parts, relatives, belongings); «перш ніж» where 'before' means "
    "«щоб не / інакше»; «але» without contrast and «і» where 'and' means «але»; needless «можете / може»; "
    "a noun repeated by a pronoun; participial phrases with a different subject; literal English idioms; "
    "detached modifiers that imply a cause; «речі» for 'things'. Return ONLY JSON with the items you changed: "
    "{\"<id>\": \"new Ukrainian text\"}; omit items that are already fine.")
REVIEW_INSTRUCTIONS = (
    "You are the terminology editor for the Ukrainian translation of Heart of the Machine. For every item, "
    "'uk' is the current translation of 'en'; 'terms' lists glossary terms that must be used (inflect "
    "as needed). Fix terminology so it matches the glossary, and fix obvious errors, but otherwise keep the "
    "translation. Keep {placeholders}, <tags>, line breaks and leading/trailing spaces. Return ONLY JSON "
    "{\"<id>\": \"corrected Ukrainian text\"} for every item.")


def record_of(key):
    p = key.split("|")
    return key if len(p) == 1 else "|".join(p[:2])


def load_manifest(round_dir):
    path = os.path.join(round_dir, "manifest.json")
    if os.path.exists(path):
        return json.load(open(path, encoding="utf-8"))
    return None


def save_manifest(round_dir, man):
    with open(os.path.join(round_dir, "manifest.json"), "w", encoding="utf-8") as f:
        json.dump(man, f, ensure_ascii=False, indent=1)


def all_rounds():
    if not os.path.isdir(BATCH_ROOT):
        return []
    res = []
    for d in sorted(os.listdir(BATCH_ROOT)):
        rd = os.path.join(BATCH_ROOT, d)
        man = load_manifest(rd) if os.path.isdir(rd) else None
        if man:
            res.append((rd, man))
    return res


def keys_in_open_batches():
    keys = set()
    for rd, man in all_rounds():
        for bid, b in man["batches"].items():
            if b["status"] != "imported":
                data = json.load(open(os.path.join(rd, bid + ".json"), encoding="utf-8"))
                if man["type"] in ("glossary", "glossary_review"):
                    keys.update("glossary:" + it["en"] for it in data["items"])
                else:
                    keys.update(it["k"] for it in data["items"])
    return keys


def new_round_dir(prefix):
    os.makedirs(BATCH_ROOT, exist_ok=True)
    n = 1
    while os.path.exists(os.path.join(BATCH_ROOT, f"{prefix}{n:02d}")):
        n += 1
    rd = os.path.join(BATCH_ROOT, f"{prefix}{n:02d}")
    os.makedirs(os.path.join(rd, "out"))
    return rd


def split_batches(items, words_per_batch, max_items, group_of):
    batches, cur, words, last_group = [], [], 0, None
    for it in items:
        g = group_of(it)
        if cur and g != last_group and (words >= words_per_batch or len(cur) >= max_items):
            batches.append(cur)
            cur, words = [], 0
        cur.append(it)
        words += len(it["en"].split()) + len((it.get("uk") or "").split())
        last_group = g
    if cur:
        batches.append(cur)
    return batches


def cmd_export(game, args):
    open_keys = keys_in_open_batches()
    partial = False
    if args.glossary_review:
        rows = [r for r in read_glossary() if r.get(G_STATUS, "") in ("proposed", "ok")
                and r.get(G_UK, "").strip() and "glossary:" + r[G_EN] not in open_keys]
        rows.sort(key=lambda r: r[G_EN].lower())
        items = [{"en": r[G_EN], "uk": r[G_UK], "note": r.get(G_NOTE, ""), "var": r.get(G_VAR, ""),
                  "cat": r.get(G_CAT, ""), "freq": r.get(G_FREQ, ""),
                  **({"approved": True} if r.get(G_STATUS) == "ok" else {})} for r in rows]
        btype, instr, prefix = "glossary_review", GLOSSARY_REVIEW_INSTRUCTIONS, "q"
        batches = [items[i:i + 5000] for i in range(0, len(items), 5000)]
    elif args.glossary:
        rows = [r for r in read_glossary()
                if r.get(G_STATUS, "") == "candidate" and not r.get(G_UK, "").strip()
                and "glossary:" + r[G_EN] not in open_keys]
        rows.sort(key=lambda r: r[G_EN].lower())  # схожі терміни (однина/множина) поруч
        items = [{"en": r[G_EN], "cat": r.get(G_CAT, ""), "de": r.get(G_DE, ""), "ex": r.get(G_EX, "")}
                 for r in rows]
        btype, instr, prefix = "glossary", GLOSSARY_INSTRUCTIONS, "g"
        batches = [items[i:i + args.items] for i in range(0, len(items), args.items)]
    else:
        rows, _, _ = read_work(WORK_FILE)
        matcher = TermMatcher(glossary_entries())
        if args.review:
            src = terms_violations(rows, matcher)
            pick = collections.OrderedDict()
            for key, e in src:
                pick.setdefault(key, []).append(e)
            by_key = {r[COL_KEY]: r for r in rows}
            items = []
            for key, ents in pick.items():
                r = by_key[key]
                if key in open_keys:
                    continue
                items.append({"k": key, "en": r[COL_EN], "uk": r[COL_UK],
                              "terms": {e[G_EN]: glossary_hint(e) for e in ents}})
            btype, instr, prefix = "review", REVIEW_INSTRUCTIONS, "e"
        elif args.shorten:
            places = ui_places(rows)
            items = []
            for r in rows:
                uk = r.get(COL_UK, "")
                if not uk.strip() or r[COL_KEY] in open_keys:
                    continue
                place = places.get(r[COL_KEY], PLACE_TEXT)
                if too_long(place, r[COL_EN], uk):
                    e = len(visible_text(r[COL_EN]))
                    limit = max(e + 3, int(e * 1.25) + 1) if place == PLACE_BUTTON else int(e * 1.5) + 3
                    items.append({"k": r[COL_KEY], "en": r[COL_EN], "uk": uk, "place": place, "max_len": limit})
            btype, instr, prefix = "review", SHORTEN_INSTRUCTIONS, "e"
        elif args.style:
            places = ui_places(rows)
            items = []
            for r in rows:
                uk = r.get(COL_UK, "")
                if not uk.strip() or r[COL_KEY] in open_keys:
                    continue
                place = places.get(r[COL_KEY], PLACE_TEXT)
                if place == PLACE_TEXT or (place == PLACE_NAME and len(r[COL_EN].split()) >= 4):
                    items.append({"k": r[COL_KEY], "en": r[COL_EN], "uk": uk})
            btype, instr, prefix = "review", STYLE_INSTRUCTIONS, "s"
            partial = True
        elif args.fix_issues:
            items = []
            for r in rows:
                uk = r.get(COL_UK, "")
                if not uk.strip() or r[COL_KEY] in open_keys:
                    continue
                issues = check_row(r[COL_KEY], r[COL_EN], uk)
                if issues:
                    items.append({"k": r[COL_KEY], "en": r[COL_EN], "uk": uk, "issues": issues})
            btype, instr, prefix = "review", FIX_INSTRUCTIONS, "e"
        else:
            statuses = set(args.statuses.split(","))
            items = []
            for r in rows:
                if r.get(COL_STATUS, "") not in statuses or r[COL_KEY] in open_keys:
                    continue
                it = {"k": r[COL_KEY], "en": r[COL_EN]}
                if r.get(COL_NOTE):
                    it["note"] = r[COL_NOTE]
                de = r.get(ref_col("German"), "")
                if de:
                    it["de"] = de
                if r.get(COL_STATUS) == "changed" and r.get(COL_UK):
                    it["old_en"], it["old_uk"] = r.get(COL_OLD_EN, ""), r[COL_UK]
                items.append(it)
            btype, instr, prefix = "strings", STRINGS_INSTRUCTIONS, "r"
        batches = split_batches(items, args.words, args.items, lambda it: record_of(it["k"]))
    if not batches:
        print("Немає що експортувати.")
        return
    rd = new_round_dir(prefix)
    man = {"type": btype, "created": datetime.datetime.now().isoformat(" ", "seconds"),
           "game_version": game.version(), "batches": {}}
    width = 3 if len(batches) < 1000 else 4
    for i, b in enumerate(batches, 1):
        bid = f"{i:0{width}d}"
        for n, it in enumerate(b, 1):
            it["id"] = n
        data = {"batch": f"{os.path.basename(rd)}/{bid}", "type": btype, "instructions": instr}
        if partial:
            data["partial"] = True  # відповідь містить лише змінені id
        if btype == "strings":
            gl = collections.OrderedDict()
            for it in b:
                for e in matcher.find(it["en"]):
                    gl.setdefault(e[G_EN], glossary_hint(e))
            data["glossary"] = gl
        data["items"] = [dict([("id", it["id"])] + [(k, v) for k, v in it.items() if k != "id"]) for it in b]
        with open(os.path.join(rd, bid + ".json"), "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=1)
        man["batches"][bid] = {"items": len(b), "words": sum(len(it["en"].split()) for it in b),
                               "first": b[0].get("k", b[0]["en"]), "last": b[-1].get("k", b[-1]["en"]),
                               "status": "pending", "by": ""}
    save_manifest(rd, man)
    total_words = sum(b["words"] for b in man["batches"].values())
    print(f"Раунд {rd}\n  тип: {btype}, батчів: {len(batches)}, елементів: {sum(len(b) for b in batches)}, "
          f"англ. слів: {total_words}")
    print(f"  відповіді ШІ кладіть у {os.path.join(rd, 'out')}\\<номер>.json (напр. out\\{'1'.zfill(width)}.json)")


def find_round(name):
    rounds = all_rounds()
    if name:
        rounds = [(rd, m) for rd, m in rounds if os.path.basename(rd) == name]
        if not rounds:
            sys.exit(f"Раунд {name} не знайдено.")
        return rounds[0]
    open_r = [(rd, m) for rd, m in rounds if any(b["status"] == "pending" for b in m["batches"].values())]
    if not open_r:
        sys.exit("Немає раундів з вільними батчами. Створіть: python hotm_uk.py export")
    return open_r[-1]


def cmd_claim(game, args):
    rd, man = find_round(args.round)
    free = [bid for bid, b in sorted(man["batches"].items()) if b["status"] == "pending"]
    if args.from_end:
        free.reverse()
    if args.ids:
        want = [x.strip().zfill(len(free[0]) if free else 3) for x in args.ids.split(",")]
        pick = [b for b in want if b in free]
    else:
        pick = free[:args.count]
    now = datetime.datetime.now().isoformat(" ", "seconds")
    matcher = TermMatcher(glossary_entries()) if man["type"] == "strings" else None
    for bid in pick:
        man["batches"][bid].update(status="claimed", by=args.name, claimed=now)
        if matcher:  # свіжі терміни з глосарію на момент взяття батчу
            path = os.path.join(rd, bid + ".json")
            data = json.load(open(path, encoding="utf-8"))
            gl = collections.OrderedDict()
            for it in data["items"]:
                for e in matcher.find(it["en"]):
                    gl.setdefault(e[G_EN], glossary_hint(e))
            data["glossary"] = gl
            items = data.pop("items")
            data["items"] = items  # items останніми — зручніше читати
            with open(path, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=1)
    save_manifest(rd, man)
    print(f"{args.name} бере {len(pick)} батч(ів) з раунду {os.path.basename(rd)}:")
    for bid in pick:
        print(f"  {os.path.join(rd, bid + '.json')}  ->  {os.path.join(rd, 'out', bid + '.json')}")


def cmd_release(game, args):
    """Повернути у вільні взяті батчі, для яких ще немає відповіді (напр. ШІ обірвався на ліміті)."""
    n = 0
    for rd, man in all_rounds():
        changed = False
        for bid, b in man["batches"].items():
            if b["status"] != "claimed" or (args.name and b.get("by") != args.name):
                continue
            if os.path.exists(os.path.join(rd, "out", bid + ".json")):
                continue  # відповідь уже є — її забере import
            b.update(status="pending", by="", claimed="")
            changed = True
            n += 1
            print(f"  {os.path.basename(rd)}/{bid} знову вільний")
        if changed:
            save_manifest(rd, man)
    print(f"Звільнено батчів: {n}")


def cmd_batches(game, args):
    rounds = all_rounds()
    if not rounds:
        print("Батчів ще немає.")
        return
    for rd, man in rounds:
        c = collections.Counter(b["status"] for b in man["batches"].values())
        by = collections.Counter(b["by"] for b in man["batches"].values() if b["by"])
        pending = sorted(bid for bid, b in man["batches"].items() if b["status"] == "pending")
        line = (f"{os.path.basename(rd):6} {man['type']:9} батчів {len(man['batches']):4}: "
                f"імпортовано {c['imported']}, взято {c['claimed']}, вільних {c['pending']}")
        if pending:
            line += f" (з початку: {pending[0]}, з кінця: {pending[-1]})"
        if by:
            line += "; " + ", ".join(f"{k}: {v}" for k, v in by.items())
        print(line)
        outs = [f for f in glob.glob(os.path.join(rd, "out", "*.json"))
                if man["batches"].get(os.path.splitext(os.path.basename(f))[0], {}).get("status") != "imported"]
        if outs:
            print(f"       готові до імпорту: {len(outs)}")


def parse_llm_json(text):
    text = text.strip().lstrip("﻿")
    text = re.sub(r"^```[a-zA-Z]*\s*", "", text)
    text = re.sub(r"\s*```\s*$", "", text)
    i, j = text.find("{"), text.rfind("}")
    if i < 0 or j < 0:
        raise ValueError("у відповіді немає JSON-об'єкта")
    return json.loads(text[i:j + 1])


def cmd_import(game, args):
    jobs = []  # (round_dir, manifest, bid, result_path)
    if args.files:
        for f in args.files:
            rd = os.path.dirname(os.path.dirname(os.path.abspath(f)))
            man = load_manifest(rd)
            bid = os.path.splitext(os.path.basename(f))[0]
            if not man or bid not in man["batches"]:
                print(f"  ! {f}: не знаю, до якого батчу це належить (очікую batches/<раунд>/out/<номер>.json)")
                continue
            jobs.append((rd, man, bid, f))
    else:
        for rd, man in all_rounds():
            for f in sorted(glob.glob(os.path.join(rd, "out", "*.json"))):
                bid = os.path.splitext(os.path.basename(f))[0]
                if bid in man["batches"] and (man["batches"][bid]["status"] != "imported" or args.force):
                    jobs.append((rd, man, bid, f))
    if not jobs:
        print("Немає нових відповідей для імпорту.")
        return

    rows, removed, info = read_work(WORK_FILE)
    by_key = {r[COL_KEY]: r for r in rows}
    gloss = read_glossary()
    g_by = {r[G_EN]: r for r in gloss}
    report = []
    touched_work = touched_gloss = False
    totals = collections.Counter()
    for rd, man, bid, path in jobs:
        batch = json.load(open(os.path.join(rd, bid + ".json"), encoding="utf-8"))
        try:
            ans = parse_llm_json(open(path, encoding="utf-8-sig").read())
        except Exception as ex:
            print(f"  ! {path}: не вдалося прочитати JSON: {ex}")
            continue
        who = args.by or man["batches"][bid].get("by") or "ai"
        st = collections.Counter()
        for it in batch["items"]:
            val = ans.get(str(it["id"]))
            if val is None and batch.get("partial"):
                st["unchanged"] += 1
                continue
            if batch["type"] == "glossary_review":
                r = g_by.get(it["en"])
                if val is None or r is None:
                    continue
                if r.get(G_STATUS) == "ok" and not args.force:
                    st["skipped_approved"] += 1
                    continue
                if isinstance(val, str):
                    val = {"uk": val}
                if val.get("skip"):
                    r[G_STATUS] = "skip"
                    st["to_skip"] += 1
                else:
                    for fld, col in (("uk", G_UK), ("note", G_NOTE), ("var", G_VAR)):
                        if fld in val:
                            r[col] = norm_text(val[fld]).strip()
                    st["changed"] += 1
                touched_gloss = True
                continue
            if val is None:
                st["missing"] += 1
                report.append((batch["batch"], it.get("k", it["en"]), "немає перекладу у відповіді", it["en"], ""))
                continue
            if batch["type"] == "glossary":
                r = g_by.get(it["en"])
                if r is None:
                    continue
                if r.get(G_STATUS) in ("ok",) and not args.force:
                    st["skipped"] += 1
                    continue
                if isinstance(val, dict) and val.get("skip"):
                    r[G_STATUS] = "skip"
                else:
                    if isinstance(val, str):
                        val = {"uk": val}
                    r[G_UK] = norm_text(val.get("uk", "")).strip()
                    if val.get("note"):
                        r[G_NOTE] = norm_text(val["note"]).strip()
                    if val.get("var"):
                        r[G_VAR] = norm_text(val["var"]).strip()
                    r[G_STATUS] = "proposed"
                st["ok"] += 1
                touched_gloss = True
                continue
            if isinstance(val, dict):
                val = val.get("uk", "")
            val = norm_text(val)
            key = it["k"]
            r = by_key.get(key)
            if r is None:
                st["gone"] += 1
                continue
            if r[COL_EN] != it["en"]:
                st["stale"] += 1
                report.append((batch["batch"], key, "англійський текст змінився після експорту — пропущено",
                               r[COL_EN], val))
                continue
            if batch["type"] == "strings" and r.get(COL_STATUS) in ("ok", "mt", "auto") \
                    and r.get(COL_UK, "").strip() and not args.force:
                st["skipped"] += 1
                continue
            if batch["type"] == "review" and r.get(COL_UK, "") != it.get("uk", ""):
                st["skipped"] += 1  # рядок змінили вручну після експорту
                continue
            if not val.strip():
                st["missing"] += 1
                continue
            if sorted(PLACEHOLDER_RE.findall(it["en"])) != sorted(PLACEHOLDER_RE.findall(val)):
                st["rejected"] += 1
                report.append((batch["batch"], key, "ВІДХИЛЕНО: плейсхолдери {..} не збігаються", it["en"], val))
                continue
            for issue in check_row(key, it["en"], val):
                report.append((batch["batch"], key, issue, it["en"], val))
                st["warn"] += 1
            if batch["type"] == "review":
                if val != r[COL_UK]:
                    r[COL_UK] = val
                    if r.get(COL_STATUS) == "ok":
                        r[COL_STATUS] = "mt"
                    r[COL_BY] = (r.get(COL_BY, "") + f"+{who}").strip("+")
                    st["ok"] += 1
                else:
                    st["unchanged"] += 1
                if it.get("terms"):  # редактор підтвердив цей переклад — не показувати в terms знову
                    terms_ignore_add(key, val)
            else:
                r[COL_UK], r[COL_STATUS], r[COL_BY], r[COL_OLD_EN] = val, "mt", who, ""
                st["ok"] += 1
            touched_work = True
        man["batches"][bid].update(status="imported", by=who,
                                   imported=datetime.datetime.now().isoformat(" ", "seconds"),
                                   result=dict(st))
        save_manifest(rd, man)
        totals.update(st)
        print(f"  {batch['batch']}: " + ", ".join(f"{k} {v}" for k, v in st.items()))
    if touched_work:
        backup_work()
        save_work(WORK_FILE, rows, removed, info)
    if touched_gloss:
        save_glossary(gloss)
    if report:
        stamp = datetime.datetime.now().strftime("%Y-%m-%d_%H%M%S")
        rp = os.path.join(BATCH_ROOT, f"import_{stamp}.txt")
        with open(rp, "w", encoding="utf-8") as f:
            for b, key, issue, en, uk in report:
                f.write(f"[{b}] {key}\n  {issue}\n  EN: {en!r}\n  UK: {uk!r}\n\n")
        print(f"Зауваження: {rp}")
    print("Разом: " + ", ".join(f"{k} {v}" for k, v in totals.items()))
    if touched_work:
        print_status(rows)


def cmd_replace(game, args):
    """Масова заміна в колонці «Українська» (редактура термінів)."""
    rows, removed, info = read_work(WORK_FILE)
    flags = 0 if args.case else re.I
    rx = re.compile(args.pattern if args.regex else re.escape(args.pattern), flags)
    changed = []
    for r in rows:
        uk = r.get(COL_UK, "")
        if not uk:
            continue
        new = rx.sub(args.repl if args.regex else (lambda m: args.repl), uk)
        if new != uk:
            changed.append((r, uk, new))
    for r, old, new in changed[:args.show]:
        print(f"{r[COL_KEY]}\n  - {old[:160]!r}\n  + {new[:160]!r}")
    if len(changed) > args.show:
        print(f"... і ще {len(changed) - args.show}")
    print(f"Рядків до зміни: {len(changed)}")
    if args.dry_run or not changed:
        return
    for r, old, new in changed:
        r[COL_UK] = new
    backup_work()
    save_work(WORK_FILE, rows, removed, info)
    print("Збережено.")


def shape_mismatch(en, uk):
    """Скільки «формальних» ознак не збігається між оригіналом і перекладом (для пошуку зсувів)."""
    score, why = 0, []
    if sorted(PLACEHOLDER_RE.findall(en)) != sorted(PLACEHOLDER_RE.findall(uk)):
        score += 2; why.append("плейсхолдери")
    if sorted(TAG_RE.findall(en)) != sorted(TAG_RE.findall(uk)):
        score += 1; why.append("теги")
    if sorted(re.findall(r"\d+", TAG_RE.sub("", en))) != sorted(re.findall(r"\d+", TAG_RE.sub("", uk))):
        score += 1; why.append("числа")
    if en.count("\n") != uk.count("\n"):
        score += 1; why.append("рядки")
    e, u = visible_text(en), visible_text(uk)
    if len(e) >= 12:
        ratio = len(u) / max(1, len(e))
        if ratio < 0.45 or ratio > 2.6:
            score += 2; why.append(f"довжина x{ratio:.1f}")
    end_e = e[-1:] if e[-1:] in "?!:." else ""
    end_u = u[-1:] if u[-1:] in "?!:.…" else ""
    if end_e in ("?", "!") and end_u != end_e:
        score += 1; why.append("кінцевий знак")
    if len(e) >= 6 and e.isupper() != u.isupper() and re.search(r"[A-Z]{3}", e):
        score += 1; why.append("капслок")
    return score, why


def cmd_audit(game, args):
    """Перевірка імпортованих відповідей ШІ: набір id, зсуви, чужий текст."""
    rows, _, _ = read_work(WORK_FILE)
    by_key = {r[COL_KEY]: r for r in rows}
    uk_index = collections.defaultdict(set)
    for r in rows:
        if len(r.get(COL_UK, "")) > 25:
            uk_index[r[COL_UK]].add(r[COL_EN])
    total_flags = 0
    for rd, man in all_rounds():
        if man["type"] not in ("strings",):
            continue
        for bid, b in sorted(man["batches"].items()):
            if args.by and b.get("by") != args.by:
                continue
            out = os.path.join(rd, "out", bid + ".json")
            if not os.path.exists(out):
                continue
            batch = json.load(open(os.path.join(rd, bid + ".json"), encoding="utf-8"))
            try:
                ans = parse_llm_json(open(out, encoding="utf-8-sig").read())
            except Exception as ex:
                print(f"{os.path.basename(rd)}/{bid} [{b.get('by')}]: НЕВАЛІДНИЙ JSON ({ex})")
                total_flags += 1
                continue
            ids = [str(it["id"]) for it in batch["items"]]
            missing = [i for i in ids if i not in ans]
            extra = [i for i in ans if i not in set(ids)]
            scores, bad, dup = [], [], []
            for it in batch["items"]:
                v = ans.get(str(it["id"]))
                if v is None:
                    continue
                v = norm_text(v if isinstance(v, str) else v.get("uk", ""))
                s, why = shape_mismatch(it["en"], v)
                scores.append(s)
                if s >= 2:
                    bad.append((it["id"], it["k"], why, it["en"], v))
                others = uk_index.get(v, set()) - {it["en"]}
                if others:
                    dup.append((it["id"], it["k"], next(iter(others))))
            # зсув: кілька поспіль підозрілих рядків
            run = best = 0
            for s in scores:
                run = run + 1 if s >= 2 else 0
                best = max(best, run)
            flag = missing or extra or best >= 3 or len(bad) > max(3, len(ids) * 0.05) or len(dup) >= 3
            if flag or args.verbose:
                total_flags += bool(flag)
                print(f"{os.path.basename(rd)}/{bid} [{b.get('by')}]: id {len(ans)}/{len(ids)}"
                      + (f", ПРОПУЩЕНО {len(missing)}" if missing else "")
                      + (f", ЗАЙВІ {len(extra)}" if extra else "")
                      + f", підозрілих {len(bad)} (найдовша серія {best}), чужий текст {len(dup)}"
                      + ("  <<< ПЕРЕВІРТЕ" if flag else ""))
                for i, k, why, en, v in bad[:args.show]:
                    print(f"     id {i} {k[:50]}: {', '.join(why)}\n        EN: {en[:90]!r}\n        UK: {v[:90]!r}")
                for i, k, other in dup[:args.show]:
                    print(f"     id {i} {k[:50]}: такий самий переклад має інший рядок: {other[:70]!r}")
    print("Підозрілих батчів:", total_flags)


def uk_count(n, one, few, many):
    """20 694 рядки / 21 рядок / 25 рядків."""
    form = one if n % 10 == 1 and n % 100 != 11 else few if 2 <= n % 10 <= 4 and not 12 <= n % 100 <= 14 else many
    return f"{n:,}".replace(",", " ") + " " + form


def read_changelog(tpl):
    """dist_template/changelog.txt → [(версія, дата, версія гри, [пункти])], найновіша перша."""
    path = os.path.join(tpl, "changelog.txt")
    if not os.path.exists(path):
        return []
    entries = []
    for line in open(path, encoding="utf-8"):
        line = line.rstrip()
        if not line or line.startswith("#!") or (line.startswith("#") and not line.startswith("##")):
            continue
        if line.startswith("## "):
            parts = [p.strip() for p in line[3:].split("|")]
            parts += [""] * (3 - len(parts))
            entries.append((parts[0], parts[1], parts[2], []))
        elif line.startswith("- ") and entries:
            entries[-1][3].append(line[2:].strip())
    return entries


def fmt_date_uk(iso):
    try:
        return datetime.date.fromisoformat(iso).strftime("%d.%m.%Y")
    except ValueError:
        return iso


def changelog_bbcode(entries):
    out = []
    for ver, date, gv, items in entries:
        head = f"[b]{ver}[/b]" + (f" — {fmt_date_uk(date)}" if date else "") + (f", гра {gv}" if gv else "")
        out.append(head + "\n[list]\n" + "".join(f"[*]{it}\n" for it in items) + "[/list]")
    return "\n".join(out)


def changelog_plain(entries):
    out = []
    for ver, date, gv, items in entries:
        head = ver + (f" ({fmt_date_uk(date)}" + (f", гра {gv}" if gv else "") + ")" if date else "")
        out.append(head + "\n" + "".join(f"  • {it}\n" for it in items))
    return "\n".join(out).rstrip()


def cmd_dist(game, args):
    """Зібрати архів українізатора для інших гравців (dist/)."""
    tpl = os.path.join(HERE, "dist_template")
    changelog = read_changelog(tpl)
    info_path = os.path.join(tpl, "info.json")
    info = json.load(open(info_path, encoding="utf-8")) if os.path.exists(info_path) else {}
    rows, _, _ = read_work(WORK_FILE)
    pairs = [(r[COL_KEY], r[COL_UK]) for r in rows if r.get(COL_UK, "").strip()]
    version = game.version()
    today = datetime.date.today().isoformat()
    release = args.release or today
    subs = {
        "{{RELEASE}}": release, "{{GAME_VERSION}}": version, "{{DATE}}": today,
        "{{COUNT}}": uk_count(len(pairs), "рядок", "рядки", "рядків") + (
            f" з {len(rows):,}".replace(",", " ") if len(pairs) < len(rows) else ""),
        "{{AUTHOR}}": info.get("author") or "(вкажіть автора в dist_template/info.json)",
        "{{CONTACT}}": info.get("contact") or "коментарі під посібником у Steam",
        "{{DOWNLOAD_URL}}": info.get("download_url") or "ПОСИЛАННЯ_НА_GOOGLE_DRIVE",
        "{{BANNER_URL}}": info.get("banner_url") or "ПОСИЛАННЯ_НА_БАНЕР",
    }
    if changelog and changelog[0][0] != release:
        print(f"УВАГА: у dist_template/changelog.txt найновіша версія {changelog[0][0]}, а збираємо {release} — "
              f"допишіть запис.")

    def fill(text, plain=False):
        text = text.replace("{{CHANGELOG}}", changelog_plain(changelog) if plain else changelog_bbcode(changelog))
        if not info.get("author"):  # автора не вказано — прибираємо рядки/розділ про автора
            text = re.sub(r"\[h2\]Автори\[/h2\]\s*\{\{AUTHOR\}\}\s*", "", text)
            text = "\n".join(l for l in text.split("\n") if "{{AUTHOR}}" not in l)
        if not info.get("banner_url"):
            text = text.replace("[img]{{BANNER_URL}}[/img]",
                                "[ТУТ ВСТАВТЕ БАНЕР: редактор посібника → «Зображення» → завантажте banner_uk.png "
                                "→ клацніть на нього, щоб вставити; цей рядок видаліть]")
        for k, v in subs.items():
            text = text.replace(k, v)
        return text

    name = f"HotM_Ukrainian_{release}"
    out_root = os.path.join(HERE, "dist")
    out = os.path.join(out_root, name)
    if os.path.exists(out):
        shutil.rmtree(out)
    os.makedirs(out)
    write_arcen_xlsx(game, pairs + link_duplicates(game, rows, pairs), os.path.join(out, "Ukrainian.xlsx"))
    shutil.copy2(LOCTYPE_BLOCK_FILE, os.path.join(out, "localization_type_uk.xml"))
    for fn in ("install.ps1", "uninstall.ps1", "Встановити.bat", "Видалити.bat", "install.sh", "uninstall.sh"):
        shutil.copy2(os.path.join(tpl, fn), os.path.join(out, fn))
    readme = fill(open(os.path.join(tpl, "ЧИТАЙМЕНЕ.txt"), encoding="utf-8").read(), plain=True)
    with open(os.path.join(out, "ЧИТАЙМЕНЕ.txt"), "w", encoding="utf-8-sig", newline="\r\n") as f:
        f.write(readme)
    guide = fill(open(os.path.join(tpl, "steam_guide.txt"), encoding="utf-8").read())
    guide_path = os.path.join(out_root, f"steam_guide_{release}.txt")
    with open(guide_path, "w", encoding="utf-8-sig", newline="\r\n") as f:
        f.write(guide)
    note_path = None
    if changelog:
        ver, date, gv, items = changelog[0]
        note = (f"Оновлення {ver}" + (f" (гра {gv})" if gv else "") + ":\n"
                + "".join(f"- {it}\n" for it in items))
        note_path = os.path.join(out_root, f"steam_changenote_{release}.txt")
        with open(note_path, "w", encoding="utf-8-sig", newline="\r\n") as f:
            f.write(note)
    for img in ("banner_uk.png", "capsule_uk.png", "capsule_uk_alt.png", "preview_uk.png"):
        if os.path.exists(os.path.join(tpl, img)):
            shutil.copy2(os.path.join(tpl, img), os.path.join(out_root, img))
    zip_path = os.path.join(out_root, name + ".zip")
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as z:
        for fn in sorted(os.listdir(out)):
            full = os.path.join(out, fn)
            zi = zipfile.ZipInfo(f"{name}/{fn}", date_time=datetime.datetime.now().timetuple()[:6])
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = (0o755 if fn.endswith(".sh") else 0o644) << 16
            z.writestr(zi, open(full, "rb").read())
    print(f"Архів: {zip_path} ({os.path.getsize(zip_path) // 1024} КБ)")
    print(f"Текст гайду для Steam: {guide_path}")
    if note_path:
        print(f"Примітка до змін для Steam: {note_path}")
    print(f"Рядків у перекладі: {len(pairs)} з {len(rows)}; версія гри {version}")


def cmd_check(game, args):
    """Перевірити відповідь ШІ без імпорту."""
    ans_path = args.answer
    if not ans_path:
        ans_path = os.path.join(os.path.dirname(os.path.abspath(args.batch)), "out", os.path.basename(args.batch))
    batch = json.load(open(args.batch, encoding="utf-8"))
    try:
        ans = parse_llm_json(open(ans_path, encoding="utf-8-sig").read())
    except Exception as ex:
        sys.exit(f"НЕВАЛІДНИЙ JSON у {ans_path}: {ex}")
    ids = {str(it["id"]) for it in batch["items"]}
    problems, notes = [], []
    extra = sorted(set(ans) - ids, key=lambda x: int(x) if x.isdigit() else 0)
    if extra:
        problems.append(f"зайві id: {', '.join(extra[:20])}")
    for it in batch["items"]:
        v = ans.get(str(it["id"]))
        if batch["type"] in ("glossary_review",):
            continue
        if v is None:
            if not batch.get("partial"):
                problems.append(f"id {it['id']}: немає перекладу")
            continue
        if batch["type"] == "glossary":
            continue
        if isinstance(v, dict):
            v = v.get("uk", "")
        v = norm_text(v)
        if not v.strip():
            problems.append(f"id {it['id']}: порожній переклад")
            continue
        if sorted(PLACEHOLDER_RE.findall(it["en"])) != sorted(PLACEHOLDER_RE.findall(v)):
            problems.append(f"id {it['id']}: плейсхолдери {{..}} не збігаються (буде ВІДХИЛЕНО)")
        for issue in check_row(it.get("k", ""), it["en"], v):
            if not issue.startswith("плейсхолдери"):
                problems.append(f"id {it['id']}: {issue}")
        for si in suspect_issues(v):
            notes.append(f"id {it['id']}: можливий русизм {si}")
        if re.search(r"[A-Za-z]{4,}", PLACEHOLDER_RE.sub("", TAG_RE.sub("", v))) and \
                v.strip() == it["en"].strip():
            notes.append(f"id {it['id']}: переклад ідентичний англійському (нормально для латинських назв, "
                         f"брендів, абревіатур — інакше перекладіть)")
    print(f"{args.batch}: {len(batch['items'])} елементів, відповідей {len(ans)}")
    for p_ in problems:
        print("  " + p_)
    for n_ in notes:
        print("  примітка: " + n_)
    print("OK" if not problems else f"Проблем: {len(problems)}")


# ---------------------------------------------------------------- редакторська перевірка термінів

TERMS_IGNORE_FILE = os.path.join(HERE, "terms_ignore.json")
_TERMS_IGNORE = None


def terms_ignore():
    global _TERMS_IGNORE
    if _TERMS_IGNORE is None:
        _TERMS_IGNORE = json.load(open(TERMS_IGNORE_FILE, encoding="utf-8")) \
            if os.path.exists(TERMS_IGNORE_FILE) else {}
    return _TERMS_IGNORE


def terms_ignore_add(key, uk):
    d = terms_ignore()
    d[key] = uk
    with open(TERMS_IGNORE_FILE, "w", encoding="utf-8") as f:
        json.dump(d, f, ensure_ascii=False, indent=0)


def terms_violations(rows, matcher):
    res = []
    ignore = terms_ignore()
    for r in rows:
        uk = r.get(COL_UK, "")
        if not uk.strip():
            continue
        if ignore.get(r[COL_KEY]) == uk:  # редактор уже підтвердив саме цей текст
            continue
        for e in matcher.find(r[COL_EN]):
            if not term_ok(e, uk):
                res.append((r[COL_KEY], e))
    return res


def cmd_terms(game, args):
    entries = glossary_entries()
    if not entries:
        sys.exit("Глосарій порожній: python hotm_uk.py glossary, потім заповніть колонку «Українська».")
    rows, _, _ = read_work(WORK_FILE)
    matcher = TermMatcher(entries)
    by_key = {r[COL_KEY]: r for r in rows}
    viol = terms_violations(rows, matcher)
    uses = collections.Counter()
    for r in rows:
        if r.get(COL_UK, "").strip():
            for e in matcher.find(r[COL_EN]):
                uses[e[G_EN]] += 1
    per_term = collections.Counter(e[G_EN] for _, e in viol)
    term_rows = [{"Термін": t, "Має бути": glossary_hint(matcher.get(t)), "Порушень": n,
                  "Вживань у перекладених": uses[t]} for t, n in per_term.most_common()]
    line_rows = []
    for key, e in viol:
        r = by_key[key]
        line_rows.append({COL_KEY: key, "Термін": e[G_EN], "Має бути": glossary_hint(e),
                          COL_EN: r[COL_EN], COL_UK: r[COL_UK], COL_STATUS: r.get(COL_STATUS, ""),
                          COL_BY: r.get(COL_BY, "")})
    path = os.path.join(HERE, "terms_report.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "По термінах"
    write_table(ws, ["Термін", "Має бути", "Порушень", "Вживань у перекладених"], term_rows, [30, 50, 11, 12],
                wrap_cols={"Має бути"})
    write_table(wb.create_sheet("Рядки"), [COL_KEY, "Термін", "Має бути", COL_EN, COL_UK, COL_STATUS, COL_BY],
                line_rows, [36, 22, 30, 60, 60, 9, 12], wrap_cols={COL_EN, COL_UK, "Має бути"})
    try:
        wb.save(path)
    except PermissionError:
        sys.exit(f"Закрийте {path} в Excel.")
    print(f"Перевірено {sum(1 for r in rows if r.get(COL_UK, '').strip())} перекладених рядків, "
          f"термінів у глосарії: {len(entries)}")
    print(f"Можливих порушень: {len(viol)} у {len({k for k, _ in viol})} рядках -> {path}")
    print("Виправити автоматично редактором-ШІ: python hotm_uk.py export --review")


def print_status(rows):
    c = collections.Counter(r.get(COL_STATUS, "") for r in rows)
    total = len(rows)
    done = c["ok"]
    words_left = sum(len(r.get(COL_EN, "").split()) for r in rows if r.get(COL_STATUS) != "ok")
    print(f"Статус: ok {done}/{total} ({100 * done / max(1, total):.1f}%), mt {c['mt']}, new {c['new']}, "
          f"changed {c['changed']}, auto {c['auto']}; ще ~{words_left} англ. слів")


def cmd_status(game, args):
    rows, removed, info = read_work(WORK_FILE)
    if not rows:
        sys.exit(f"Немає {WORK_FILE}. Спершу запустіть:  python hotm_uk.py sync")
    print(f"Версія гри при останній синхронізації: {info.get('Версія гри', '?')} "
          f"(збірка {info.get('Збірка Steam') or '?'}), поточна: {game.version()} (збірка {game.build_id() or '?'})")
    print_status(rows)


def cmd_dump(game, args):
    rows, removed, info = read_work(WORK_FILE)
    if not rows:
        sys.exit(f"Немає {WORK_FILE}. Спершу запустіть:  python hotm_uk.py sync")
    dump_work(rows, removed, info)
    dump_glossary(read_glossary())
    print(f"Текстову копію оновлено: {TEXT_DIR}")


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--game", help="тека Heart of the Machine (типово — батьківська тека скрипта)")
    sub = p.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("sync", help="оновити робочий файл з англійського тексту гри")
    s.add_argument("--dry-run", action="store_true", help="лише показати кількість змін")
    b = sub.add_parser("build", help="зібрати Ukrainian.xlsx для гри")
    b.add_argument("--skip-changed", action="store_true",
                   help="не включати рядки зі статусом changed (застарілий переклад)")
    sub.add_parser("status", help="статистика перекладу")
    sub.add_parser("dump", help="оновити текстову копію text/*.json з uk_translation.xlsx і glossary.xlsx")
    sub.add_parser("pseudo", help="тестовий Ukrainian.xlsx для перевірки шрифтів")
    sub.add_parser("glossary", help="оновити glossary.xlsx кандидатами в терміни")
    e = sub.add_parser("export", help="нарізати JSON-батчі для ШІ")
    e.add_argument("--glossary", action="store_true", help="батчі з неперекладених термінів глосарію")
    e.add_argument("--review", action="store_true", help="батчі для редактора: рядки з порушеннями глосарію")
    e.add_argument("--shorten", action="store_true",
                   help="батчі для UI-редактора: короткі написи, помітно довші за англійські")
    e.add_argument("--style", action="store_true",
                   help="батчі для літературного редактора: природність мови, русизми, кальки (відповідь — лише змінені id)")
    e.add_argument("--fix-issues", action="store_true",
                   help="батчі для технічного редактора: рядки з проблемами (переноси, теги, пробіли, заборонені слова)")
    e.add_argument("--glossary-review", action="store_true",
                   help="весь глосарій одним батчем для редактора термінів (узгодженість, skip)")
    e.add_argument("--statuses", default="new,changed", help="які статуси експортувати (типово new,changed)")
    e.add_argument("--words", type=int, default=2000, help="англ. слів на батч (типово 2000)")
    e.add_argument("--items", type=int, default=200, help="макс. елементів на батч (типово 200)")
    c = sub.add_parser("claim", help="взяти вільні батчі")
    c.add_argument("name", help="хто перекладає: claude, gemini, ...")
    c.add_argument("count", type=int, nargs="?", default=1)
    c.add_argument("--ids", help="конкретні номери батчів через кому, напр. 001,047")
    c.add_argument("--from-end", action="store_true", help="брати з кінця")
    c.add_argument("--round", help="назва раунду (типово останній з вільними батчами)")
    sub.add_parser("batches", help="стан батчів")
    rl = sub.add_parser("release", help="звільнити взяті батчі без відповіді")
    rl.add_argument("name", nargs="?", help="лише батчі цього перекладача (claude, gemini, ...)")
    i = sub.add_parser("import", help="імпортувати відповіді ШІ")
    i.add_argument("files", nargs="*", help="файли відповідей (типово всі нові batches/*/out/*.json)")
    i.add_argument("--by", help="позначка перекладача (типово з claim)")
    i.add_argument("--force", action="store_true", help="перезаписувати вже перекладені рядки")
    sub.add_parser("terms", help="перевірка термінології за глосарієм")
    ds = sub.add_parser("dist", help="зібрати архів українізатора для інших гравців (dist/)")
    ds.add_argument("--release", help="назва/версія релізу перекладу (типово сьогоднішня дата)")
    au = sub.add_parser("audit", help="перевірити імпортовані відповіді ШІ: набір id, зсуви, чужий текст")
    au.add_argument("--by", help="лише батчі цього перекладача (gemini, claude, ...)")
    au.add_argument("--show", type=int, default=4, help="скільки прикладів показувати на батч")
    au.add_argument("--verbose", action="store_true", help="показати всі батчі, не лише підозрілі")
    rp = sub.add_parser("replace", help="масова заміна в перекладі (редактура термінів)")
    rp.add_argument("pattern")
    rp.add_argument("repl")
    rp.add_argument("--regex", action="store_true", help="pattern/repl — регулярні вирази")
    rp.add_argument("--case", action="store_true", help="враховувати регістр")
    rp.add_argument("--dry-run", action="store_true", help="лише показати")
    rp.add_argument("--show", type=int, default=30, help="скільки змін показати")
    k = sub.add_parser("check", help="перевірити відповідь ШІ без імпорту")
    k.add_argument("batch", help="файл батчу, напр. batches/r01/001.json")
    k.add_argument("answer", nargs="?", help="файл відповіді (типово batches/<раунд>/out/<той самий номер>.json)")
    args = p.parse_args()
    global GAME_ROOT
    game = Game(find_game_dir(args.game))
    GAME_ROOT = game.root
    {"sync": cmd_sync, "build": cmd_build, "status": cmd_status, "dump": cmd_dump, "pseudo": cmd_pseudo,
     "glossary": cmd_glossary, "export": cmd_export, "claim": cmd_claim, "batches": cmd_batches,
     "import": cmd_import, "terms": cmd_terms, "check": cmd_check, "replace": cmd_replace,
     "release": cmd_release, "audit": cmd_audit, "dist": cmd_dist}[args.cmd](game, args)


if __name__ == "__main__":
    main()
